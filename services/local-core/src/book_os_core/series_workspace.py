from __future__ import annotations

import hashlib
from html import unescape
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any, Literal, cast
from zipfile import BadZipFile, ZipFile

from docx import Document
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from sqlalchemy import text

from .authority import new_ulid
from .authority_types import utc_now
from .book_context import ProfileView
from .series_similarity import semantic_families, semantic_score, semantic_series_findings
from .series_workspace_base import *  # noqa: F401,F403
from .series_workspace_base import (
    SeriesBookCreateRequest,
    SeriesBookLifecycle,
    SeriesBookStatus,
    SeriesBookView,
    SeriesCreateRequest,
    SeriesMapView,
    SeriesWorkspaceGateError,
    SeriesWorkspaceService as _BaseSeriesWorkspaceService,
)

SeriesScopeRecommendation = Literal["NEW_BOOK", "EXISTING_BOOK_CHAPTER", "REVIEW"]
OverlapClassification = Literal[
    "TERM",
    "SHORT_REMINDER",
    "DEVELOPMENT",
    "NEW_CONTEXT_APPLICATION",
    "UNACCEPTABLE_DUPLICATE",
]


class SeriesBookScopeRequest(BaseModel):
    idea: str = Field(min_length=3, max_length=12000)
    reader_problem: str = Field(default="", max_length=12000)
    reader_result: str = Field(default="", max_length=12000)
    unique_mechanism: str = Field(default="", max_length=12000)


class SeriesBookScopeAssessment(BaseModel):
    series_profile_id: str
    recommendation: SeriesScopeRecommendation
    confidence: float = Field(ge=0, le=1)
    candidate_book_id: str | None = None
    candidate_book_title: str | None = None
    evidence: dict[str, Any] = Field(default_factory=dict)


class SeriesTopicOwnershipRequest(BaseModel):
    topic_label: str = Field(min_length=3, max_length=1000)
    owner_book_id: str = Field(min_length=26, max_length=26)
    reason: str = Field(min_length=8, max_length=6000)


class SeriesTopicOwnershipView(BaseModel):
    ownership_id: str
    series_profile_id: str
    topic_key: str
    topic_label: str
    owner_book_id: str
    owner_book_title: str
    reason: str
    actor: str
    supersedes_ownership_id: str | None = None
    created_at: str


class SeriesOverlapDispositionRequest(BaseModel):
    classification: OverlapClassification
    reason: str = Field(min_length=8, max_length=6000)
    topic_label: str | None = Field(default=None, min_length=3, max_length=1000)
    owner_book_id: str | None = Field(default=None, min_length=26, max_length=26)


class SeriesOverlapDispositionView(BaseModel):
    finding_id: str
    classification: OverlapClassification
    finding_status: Literal["OPEN", "RESOLVED", "ACCEPTED_EXCEPTION"]
    map: SeriesMapView


class SeriesWorkspaceService(_BaseSeriesWorkspaceService):  # type: ignore[no-redef]
    """Durable nonfiction-series governance layered on the original workspace."""

    _SERIES_REVIEW_FILE = "series-review-pending.json"
    _SEMANTIC_THESIS_BLOCK_SCORE = 0.70

    def _map_hash(self, series_profile_id: str, books: list[SeriesBookView]) -> str:
        """Hash only material series-comparison inputs, not workflow bookkeeping.

        A difference map must become stale when a Series Bible, Book Passport, current
        architecture, imported source, or archive inclusion changes. It must not become stale
        merely because an otherwise identical book advances from WRITING to EDITING/FINAL_REVIEW
        or receives a new `updated_at` timestamp.
        """
        profile = self.profiles.get_profile(series_profile_id)
        material = {
            "profile_hash": profile.content_hash,
            "books": [
                {
                    "book_id": book.book_id,
                    "title": book.title,
                    "ordinal": book.ordinal,
                    "unique_idea": book.unique_idea,
                    "reader_problem": book.reader_problem,
                    "reader_result": book.reader_result,
                    "unique_mechanism": book.unique_mechanism,
                    "excluded_topics": book.excluded_topics,
                    "source_kind": book.source_kind,
                    "origin_kind": book.origin_kind,
                    "legacy_content_allowed": book.legacy_content_allowed,
                    "passport_hash": book.passport_hash,
                    "archived": book.lifecycle == "ARCHIVED",
                }
                for book in books
            ],
            "architectures": {
                book.book_id: self._architecture_material(book.book_id) for book in books
            },
            "sources": {
                book.book_id: sorted(
                    row["content_hash"]
                    for row in (
                        []
                        if book.origin_kind == "LEGACY_TITLE_ONLY"
                        else self._source_rows(book.book_id)
                    )
                )
                for book in books
            },
        }
        return hashlib.sha256(
            json.dumps(
                material,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()

    def _review_path(self) -> Path:
        return self.data_dir / self._SERIES_REVIEW_FILE

    def _review_pending(self) -> set[str]:
        path = self._review_path()
        if not path.exists():
            return set()
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return set()
        values = payload.get("series_profile_ids", []) if isinstance(payload, dict) else []
        return {str(value) for value in values if isinstance(value, str)}

    def _write_review_pending(self, values: set[str]) -> None:
        path = self._review_path()
        temporary = path.with_suffix(".tmp")
        temporary.write_text(
            json.dumps(
                {"series_profile_ids": sorted(values)},
                ensure_ascii=False,
                sort_keys=True,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
        )
        temporary.replace(path)

    def requires_series_review(self, series_profile_id: str) -> bool:
        return series_profile_id in self._review_pending()

    def clear_series_review_requirement(self, series_profile_id: str) -> None:
        values = self._review_pending()
        if series_profile_id in values:
            values.remove(series_profile_id)
            self._write_review_pending(values)

    def create_series(self, request: SeriesCreateRequest) -> ProfileView:
        profile = super().create_series(request)
        values = self._review_pending()
        values.add(profile.profile_id)
        self._write_review_pending(values)
        return profile

    def _full_source_text(self, row: dict[str, Any]) -> str:
        candidate = self.projects.projects_dir / str(row["book_id"]) / str(row["relative_path"])
        if not candidate.is_file():
            return ""
        fmt = str(row["format"])
        try:
            payload = candidate.read_bytes()
            if fmt in {"TXT", "MARKDOWN"}:
                return payload.decode("utf-8", errors="replace")
            if fmt == "DOCX":
                document = Document(BytesIO(payload))
                paragraphs = [item.text for item in document.paragraphs]
                tables = [
                    "\n".join(
                        " | ".join(cell.text for cell in table_row.cells)
                        for table_row in table.rows
                    )
                    for table in document.tables
                ]
                return "\n".join([*paragraphs, *tables])
            if fmt == "PDF":
                return "\n".join(
                    page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages
                )
            if fmt == "EPUB":
                with ZipFile(BytesIO(payload)) as archive:
                    html_parts = [
                        archive.read(name).decode("utf-8", errors="replace")
                        for name in archive.namelist()
                        if name.casefold().endswith((".xhtml", ".html", ".htm"))
                    ]
                return unescape(re.sub(r"<[^>]+>", " ", "\n".join(html_parts)))
        except (BadZipFile, OSError, PdfReadError, UnicodeError, ValueError):
            return ""
        return ""

    def _source_rows(self, book_id: str) -> list[dict[str, Any]]:
        rows = super()._source_rows(book_id)
        for row in rows:
            if str(row.get("analysis_status")) != "PARSED":
                continue
            full_text = self._full_source_text(row)
            if not full_text.strip():
                continue
            analysis = cast(dict[str, Any], row["analysis"])
            analysis["sample"] = full_text
            analysis["whole_book_characters"] = len(full_text)
            analysis["whole_book_duplicate_scan"] = True
        return rows

    def assess_book_scope(
        self,
        series_profile_id: str,
        request: SeriesBookScopeRequest,
    ) -> SeriesBookScopeAssessment:
        candidate = " ".join(
            (
                request.idea,
                request.reader_problem,
                request.reader_result,
                request.unique_mechanism,
            )
        )
        best_book: SeriesBookView | None = None
        best_score = 0.0
        best_families: set[str] = set()
        candidate_families = semantic_families(candidate)
        for book in self.books(series_profile_id):
            if book.lifecycle == "ARCHIVED":
                continue
            existing = " ".join(
                (
                    book.unique_idea,
                    book.reader_problem,
                    book.reader_result,
                    book.unique_mechanism,
                )
            )
            score = semantic_score(candidate, existing)
            shared = candidate_families & semantic_families(existing)
            weighted = score + min(0.12, len(shared) * 0.02)
            if weighted > best_score:
                best_score = min(1.0, weighted)
                best_book = book
                best_families = shared
        if best_book is None:
            recommendation: SeriesScopeRecommendation = "NEW_BOOK"
        elif best_score >= 0.74 and len(best_families) >= 2:
            recommendation = "EXISTING_BOOK_CHAPTER"
        elif best_score >= 0.56 and best_families:
            recommendation = "REVIEW"
        else:
            recommendation = "NEW_BOOK"
        return SeriesBookScopeAssessment(
            series_profile_id=series_profile_id,
            recommendation=recommendation,
            confidence=round(best_score, 4),
            candidate_book_id=best_book.book_id if best_book is not None else None,
            candidate_book_title=best_book.title if best_book is not None else None,
            evidence={
                "shared_semantic_families": sorted(best_families),
                "comparison_basis": "SERIES_NEW_BOOK_VS_EXISTING_CHAPTER_V1",
                "provider_calls": 0,
            },
        )

    def bind_existing_book(
        self,
        series_profile_id: str,
        book_id: str,
        *,
        idea: str,
    ) -> SeriesBookView:
        project = self.projects.get_project(book_id)
        title_match = next(
            (
                item
                for item in self.books(series_profile_id)
                if item.title.strip().casefold() == project.working_title.strip().casefold()
            ),
            None,
        )
        if title_match is None:
            assessment = self.assess_book_scope(
                series_profile_id,
                SeriesBookScopeRequest(idea=idea),
            )
            if assessment.recommendation == "EXISTING_BOOK_CHAPTER":
                raise SeriesWorkspaceGateError(
                    "idea is too close to an existing series book and must first be treated as a "
                    f"possible chapter of «{assessment.candidate_book_title}»"
                )
        return super().bind_existing_book(series_profile_id, book_id, idea=idea)

    @staticmethod
    def _topic_key(topic_label: str) -> str:
        normalized = " ".join(re.findall(r"[а-яёa-z0-9]+", topic_label.casefold()))
        return hashlib.sha256(normalized.encode("utf-8")).hexdigest()

    def _anchor_book(self, series_profile_id: str) -> SeriesBookView:
        books = self.books(series_profile_id)
        if not books:
            raise SeriesWorkspaceGateError("series has no books for durable governance records")
        return books[0]

    def topic_ownerships(self, series_profile_id: str) -> list[SeriesTopicOwnershipView]:
        books = self.books(series_profile_id)
        if not books:
            return []
        anchor = books[0]
        titles = {item.book_id: item.title for item in books}
        engine = self._engine(anchor.book_id)
        try:
            with engine.connect() as connection:
                rows = list(
                    connection.execute(
                        text(
                            "SELECT * FROM series_topic_ownership "
                            "WHERE series_profile_id=:series ORDER BY created_at,ownership_id"
                        ),
                        {"series": series_profile_id},
                    ).mappings()
                )
        finally:
            engine.dispose()
        latest: dict[str, Any] = {}
        for row in rows:
            latest[str(row["topic_key"])] = row
        return sorted(
            [
                SeriesTopicOwnershipView(
                    ownership_id=str(row["ownership_id"]),
                    series_profile_id=str(row["series_profile_id"]),
                    topic_key=str(row["topic_key"]),
                    topic_label=str(row["topic_label"]),
                    owner_book_id=str(row["owner_book_id"]),
                    owner_book_title=titles.get(str(row["owner_book_id"]), "Архивная книга серии"),
                    reason=str(row["reason"]),
                    actor=str(row["actor"]),
                    supersedes_ownership_id=(
                        str(row["supersedes_ownership_id"])
                        if row["supersedes_ownership_id"] is not None
                        else None
                    ),
                    created_at=str(row["created_at"]),
                )
                for row in latest.values()
            ],
            key=lambda item: item.topic_label.casefold(),
        )

    def assign_topic_ownership(
        self,
        series_profile_id: str,
        request: SeriesTopicOwnershipRequest,
        *,
        actor: str = "HUMAN:OWNER",
    ) -> SeriesTopicOwnershipView:
        books = self.books(series_profile_id)
        owner = next(
            (
                item
                for item in books
                if item.book_id == request.owner_book_id and item.lifecycle != "ARCHIVED"
            ),
            None,
        )
        if owner is None:
            raise SeriesWorkspaceGateError(
                "topic owner must be an active book in this exact series"
            )
        anchor = self._anchor_book(series_profile_id)
        topic_label = request.topic_label.strip()
        topic_key = self._topic_key(topic_label)
        ownership_id = new_ulid()
        now = utc_now()
        engine = self._engine(anchor.book_id)
        try:
            with engine.begin() as connection:
                prior = connection.execute(
                    text(
                        "SELECT ownership_id FROM series_topic_ownership "
                        "WHERE series_profile_id=:series AND topic_key=:topic "
                        "ORDER BY created_at DESC,ownership_id DESC LIMIT 1"
                    ),
                    {"series": series_profile_id, "topic": topic_key},
                ).scalar_one_or_none()
                connection.execute(
                    text(
                        "INSERT INTO series_topic_ownership(ownership_id,series_profile_id,topic_key,"
                        "topic_label,owner_book_id,reason,actor,supersedes_ownership_id,created_at) "
                        "VALUES (:id,:series,:topic,:label,:owner,:reason,:actor,:prior,:created)"
                    ),
                    {
                        "id": ownership_id,
                        "series": series_profile_id,
                        "topic": topic_key,
                        "label": topic_label,
                        "owner": request.owner_book_id,
                        "reason": request.reason.strip(),
                        "actor": actor,
                        "prior": prior,
                        "created": now,
                    },
                )
        finally:
            engine.dispose()
        return next(
            item
            for item in self.topic_ownerships(series_profile_id)
            if item.ownership_id == ownership_id
        )

    def _append_effective_map_run(self, series_profile_id: str, map_hash: str) -> SeriesMapView:
        anchor = self._anchor_book(series_profile_id)
        now = utc_now()
        engine = self._engine(anchor.book_id)
        try:
            with engine.begin() as connection:
                rows = list(
                    connection.execute(
                        text(
                            "SELECT severity,status FROM series_similarity_findings "
                            "WHERE series_profile_id=:series AND map_hash=:hash"
                        ),
                        {"series": series_profile_id, "hash": map_hash},
                    ).mappings()
                )
                open_rows = [row for row in rows if str(row["status"]) == "OPEN"]
                status = (
                    "BLOCKING"
                    if any(str(row["severity"]) == "BLOCKING" for row in open_rows)
                    else "ATTENTION"
                    if open_rows
                    else "PASS"
                )
                connection.execute(
                    text(
                        "INSERT INTO series_map_runs(map_run_id,series_profile_id,book_id,map_hash,"
                        "status,finding_count,created_at) VALUES (:id,:series,:book,:hash,:status,"
                        ":count,:created)"
                    ),
                    {
                        "id": new_ulid(),
                        "series": series_profile_id,
                        "book": anchor.book_id,
                        "hash": map_hash,
                        "status": status,
                        "count": len(open_rows),
                        "created": now,
                    },
                )
        finally:
            engine.dispose()
        current = self.current_map(series_profile_id)
        if current is None:
            raise SeriesWorkspaceGateError("series map disappeared while recording disposition")
        return current

    def dispose_overlap(
        self,
        series_profile_id: str,
        finding_id: str,
        request: SeriesOverlapDispositionRequest,
        *,
        actor: str = "HUMAN:OWNER",
    ) -> SeriesOverlapDispositionView:
        current = self.current_map(series_profile_id)
        if current is None or not current.current:
            raise SeriesWorkspaceGateError(
                "build the current difference map before resolving overlap"
            )
        anchor = self._anchor_book(series_profile_id)
        engine = self._engine(anchor.book_id)
        try:
            with engine.connect() as connection:
                row = (
                    connection.execute(
                        text(
                            "SELECT * FROM series_similarity_findings WHERE finding_id=:finding "
                            "AND series_profile_id=:series AND map_hash=:hash"
                        ),
                        {
                            "finding": finding_id,
                            "series": series_profile_id,
                            "hash": current.map_hash,
                        },
                    )
                    .mappings()
                    .one_or_none()
                )
        finally:
            engine.dispose()
        if row is None:
            raise SeriesWorkspaceGateError("overlap finding is not part of the current map")
        if str(row["dimension"]) == "SOURCE_QUALITY":
            raise SeriesWorkspaceGateError(
                "source-quality blockers cannot be accepted as overlap exceptions"
            )
        if str(row["status"]) != "OPEN":
            return SeriesOverlapDispositionView(
                finding_id=finding_id,
                classification=request.classification,
                finding_status=cast(Any, str(row["status"])),
                map=current,
            )

        if request.classification == "DEVELOPMENT":
            if request.topic_label is None or request.owner_book_id is None:
                raise SeriesWorkspaceGateError(
                    "development overlap requires a canonical topic label and owner book"
                )
            self.assign_topic_ownership(
                series_profile_id,
                SeriesTopicOwnershipRequest(
                    topic_label=request.topic_label,
                    owner_book_id=request.owner_book_id,
                    reason=request.reason,
                ),
                actor=actor,
            )

        accepted = request.classification != "UNACCEPTABLE_DUPLICATE"
        finding_status = "ACCEPTED_EXCEPTION" if accepted else "OPEN"
        decision = "APPROVED" if accepted else "REVISE"
        input_hash = hashlib.sha256(
            json.dumps(
                {
                    "finding_id": finding_id,
                    "map_hash": current.map_hash,
                    "classification": request.classification,
                    "reason": request.reason.strip(),
                    "topic_label": request.topic_label,
                    "owner_book_id": request.owner_book_id,
                },
                ensure_ascii=False,
                sort_keys=True,
            ).encode("utf-8")
        ).hexdigest()
        engine = self._engine(anchor.book_id)
        try:
            with engine.begin() as connection:
                if accepted:
                    connection.execute(
                        text(
                            "UPDATE series_similarity_findings SET status='ACCEPTED_EXCEPTION' "
                            "WHERE finding_id=:finding AND status='OPEN'"
                        ),
                        {"finding": finding_id},
                    )
                connection.execute(
                    text(
                        "INSERT INTO series_author_decisions(decision_id,series_profile_id,book_id,"
                        "decision_kind,input_hash,decision,reason,actor,created_at) VALUES "
                        "(:id,:series,:book,'EXCEPTION',:hash,:decision,:reason,:actor,:created)"
                    ),
                    {
                        "id": new_ulid(),
                        "series": series_profile_id,
                        "book": anchor.book_id,
                        "hash": input_hash,
                        "decision": decision,
                        "reason": (
                            f"{request.classification}: {request.reason.strip()} "
                            f"[finding={finding_id}]"
                        ),
                        "actor": actor,
                        "created": utc_now(),
                    },
                )
        finally:
            engine.dispose()
        updated_map = self._append_effective_map_run(series_profile_id, current.map_hash)
        return SeriesOverlapDispositionView(
            finding_id=finding_id,
            classification=request.classification,
            finding_status=cast(Any, finding_status),
            map=updated_map,
        )

    def _derived_series_state(
        self, book: SeriesBookView
    ) -> tuple[SeriesBookStatus, SeriesBookLifecycle]:
        if book.lifecycle == "ARCHIVED":
            return "ARCHIVED", "ARCHIVED"
        if book.current_corpus_eligible and book.lifecycle == "COMPLETED":
            return "READY", "COMPLETED"
        project = self.projects.get_project(book.book_id)
        engine = self._engine(book.book_id)
        try:
            with engine.connect() as connection:
                runtime = (
                    connection.execute(
                        text(
                            "SELECT current_stage,status FROM auto_book_runtime_runs "
                            "WHERE book_id=:book ORDER BY updated_at DESC,run_id DESC LIMIT 1"
                        ),
                        {"book": book.book_id},
                    )
                    .mappings()
                    .first()
                )
                master_count = int(
                    connection.execute(
                        text("SELECT COUNT(*) FROM literary_masters WHERE book_id=:book"),
                        {"book": book.book_id},
                    ).scalar_one()
                )
        finally:
            engine.dispose()
        if master_count > 0 or (runtime is not None and str(runtime["status"]) == "PACKAGE_READY"):
            return "READY", "COMPLETED"
        if runtime is not None:
            stage = str(runtime["current_stage"])
            if stage in {"DEFINITION", "RESEARCH"}:
                return "DEFINITION", "DEFINITION"
            if stage == "ARCHITECTURE":
                return "ARCHITECTURE", "ARCHITECTURE"
            if stage in {"CHAPTER_CONTEXT", "WRITING", "CHAPTER_REVIEW", "MIDBOOK_AUDIT"}:
                return "WRITING", "WRITING"
            if stage in {
                "WHOLE_BOOK_EDIT",
                "FACT_CHECK",
                "LITERARY_EDIT",
                "VISUALS",
                "INDEPENDENT_CRITIQUE",
                "CORRECTION",
                "AUDIO_EDITORIAL",
            }:
                return "EDITING", "EDITING"
            if stage == "MASTER_AND_EXPORTS":
                return "FINAL_REVIEW", "EDITING"
        if project.architecture is not None and project.architecture.authority_status in {
            "APPROVED",
            "LOCKED",
        }:
            return "ARCHITECTURE", "ARCHITECTURE"
        if project.book_contract is not None:
            return "DEFINITION", "DEFINITION"
        if book.lifecycle == "PLANNED":
            return "IDEA", "PLANNED"
        return "DEFINITION", "DEFINITION"

    def books(self, series_profile_id: str) -> list[SeriesBookView]:
        books = super().books(series_profile_id)
        result: list[SeriesBookView] = []
        for book in books:
            status, lifecycle = self._derived_series_state(book)
            if status != book.status or lifecycle != book.lifecycle:
                now = utc_now()
                engine = self._engine(book.book_id)
                try:
                    with engine.begin() as connection:
                        connection.execute(
                            text(
                                "UPDATE series_books SET status=:status,lifecycle=:lifecycle,"
                                "updated_at=:updated WHERE series_book_id=:membership"
                            ),
                            {
                                "status": status,
                                "lifecycle": lifecycle,
                                "updated": now,
                                "membership": book.series_book_id,
                            },
                        )
                finally:
                    engine.dispose()
                # Lifecycle remains visible in the Series UI, but map freshness is based only on
                # semantic comparison inputs. Advancing a workflow stage must not invalidate a map.
                book = book.model_copy(
                    update={"status": status, "lifecycle": lifecycle, "updated_at": now}
                )
            result.append(book)
        return result

    def _semantic_findings(self, series_profile_id: str) -> list[dict[str, Any]]:
        books = self.books(series_profile_id)
        findings: list[dict[str, Any]] = []
        for index, left in enumerate(books):
            left_architecture = self._architecture_material(left.book_id)
            left_sources = (
                [] if left.origin_kind == "LEGACY_TITLE_ONLY" else self._source_rows(left.book_id)
            )
            for right in books[index + 1 :]:
                right_architecture = self._architecture_material(right.book_id)
                right_sources = (
                    []
                    if right.origin_kind == "LEGACY_TITLE_ONLY"
                    else self._source_rows(right.book_id)
                )
                for candidate in semantic_series_findings(
                    left,
                    right,
                    left_architecture=left_architecture,
                    right_architecture=right_architecture,
                    left_sources=left_sources,
                    right_sources=right_sources,
                ):
                    # A series shares vocabulary by design. A semantic-thesis hit becomes a hard
                    # blocker only at high confidence; weaker thematic similarity remains governed
                    # by the original lexical/boundary map. Structural/case/tool/template clones
                    # remain hard blockers regardless of profession or domain wording.
                    if candidate.dimension == "THESIS":
                        score = float(candidate.evidence.get("semantic_paraphrase_score", 0.0))
                        if score < self._SEMANTIC_THESIS_BLOCK_SCORE:
                            continue
                    findings.append(
                        {
                            "finding_id": new_ulid(),
                            "book_id": left.book_id,
                            "compared_book_id": right.book_id,
                            "dimension": candidate.dimension,
                            "severity": candidate.severity,
                            "evidence": candidate.evidence,
                        }
                    )
        return findings

    def analyze(self, series_profile_id: str) -> SeriesMapView:
        # Synchronize lifecycle before base map material/hash is built.
        self.books(series_profile_id)
        exact = super().analyze(series_profile_id)
        semantic = self._semantic_findings(series_profile_id)
        if not semantic:
            return exact

        books = self.books(series_profile_id)
        if not books:
            return exact
        anchor = books[0].book_id
        now = utc_now()
        engine = self._engine(anchor)
        persisted: list[dict[str, Any]] = []
        try:
            with engine.begin() as connection:
                existing_signatures = {
                    str(evidence.get("semantic_signature"))
                    for row in connection.execute(
                        text(
                            "SELECT evidence_json FROM series_similarity_findings "
                            "WHERE series_profile_id=:series AND map_hash=:hash"
                        ),
                        {"series": series_profile_id, "hash": exact.map_hash},
                    ).mappings()
                    if isinstance(
                        (evidence := json.loads(str(row["evidence_json"]))),
                        dict,
                    )
                    and evidence.get("semantic_signature")
                }
                for item in semantic:
                    signature = str(item["evidence"].get("semantic_signature", ""))
                    if signature and signature in existing_signatures:
                        continue
                    persisted_evidence = {
                        **cast(dict[str, Any], item["evidence"]),
                        "subject_book_id": item["book_id"],
                        "compared_book_id": item["compared_book_id"],
                    }
                    connection.execute(
                        text(
                            "INSERT INTO series_similarity_findings(finding_id,series_profile_id,"
                            "book_id,compared_book_id,dimension,severity,evidence_json,map_hash,status,"
                            "created_at) VALUES (:id,:series,:book,:compared,:dimension,'BLOCKING',"
                            ":evidence,:hash,'OPEN',:created)"
                        ),
                        {
                            "id": item["finding_id"],
                            "series": series_profile_id,
                            "book": anchor,
                            "compared": item["compared_book_id"],
                            "dimension": item["dimension"],
                            "evidence": json.dumps(
                                persisted_evidence,
                                ensure_ascii=False,
                                sort_keys=True,
                            ),
                            "hash": exact.map_hash,
                            "created": now,
                        },
                    )
                    persisted.append(item)
                    if signature:
                        existing_signatures.add(signature)
                total_findings = connection.execute(
                    text(
                        "SELECT COUNT(*) FROM series_similarity_findings "
                        "WHERE series_profile_id=:series AND map_hash=:hash"
                    ),
                    {"series": series_profile_id, "hash": exact.map_hash},
                ).scalar_one()
                connection.execute(
                    text(
                        "INSERT INTO series_map_runs(map_run_id,series_profile_id,book_id,map_hash,"
                        "status,finding_count,created_at) VALUES (:id,:series,:book,:hash,'BLOCKING',"
                        ":count,:created)"
                    ),
                    {
                        "id": new_ulid(),
                        "series": series_profile_id,
                        "book": anchor,
                        "hash": exact.map_hash,
                        "count": int(total_findings),
                        "created": now,
                    },
                )
        finally:
            engine.dispose()

        current = self.current_map(series_profile_id)
        if current is None:
            return SeriesMapView(
                map_hash=exact.map_hash,
                status="BLOCKING",
                findings=[*exact.findings, *persisted],
                approved=False,
                current=True,
            )
        return current

    def generation_context(self, book_id: str) -> dict[str, Any] | None:
        context = super().generation_context(book_id)
        if context is None:
            return None
        series_profile_id = str(context["series_profile_id"])
        return {
            **context,
            "topic_ownership": [
                item.model_dump(mode="json") for item in self.topic_ownerships(series_profile_id)
            ],
        }
