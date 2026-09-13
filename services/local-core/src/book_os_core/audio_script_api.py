from __future__ import annotations

import base64
import binascii
import hashlib
from html import unescape
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Callable, Literal, cast
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, Depends, HTTPException
import httpx
from pydantic import BaseModel, Field
from pypdf.errors import PdfReadError

from .audio_script import (
    AudioScriptContent,
    AudioScriptError,
    AudioScriptGateError,
    AudioScriptService,
    AudioScriptView,
)
from .auto_book_exports import (
    AutoBookExporter,
    MasterChapter,
    StructuredBookMaster,
)
from .auto_book_finalizer import AUDIO_SCRIPT_EDITOR_V1
from .auto_book_runtime import (
    AutoBookAttachment,
    AutoBookBudgetError,
    AutoBookIntent,
    AutoBookOutputSelection,
    AutoBookRuntimeError,
    AutoBookStage,
    DurableAutoBookRuntime,
)
from .model_gateway import (
    ModelGateway,
    ModelOutputError,
    ModelProviderError,
    ModelTaskRequest,
    SectionDraftOutput,
)
from .model_routing import ModelRoutingService


class ExistingAudioPrepareRequest(BaseModel):
    source_filename: str = Field(min_length=1, max_length=500)
    source_content_base64: str = Field(min_length=1, max_length=40_000_000)
    title: str = Field(min_length=1, max_length=500)
    author: str = Field(min_length=1, max_length=300)
    adaptation_mode: Literal["SOURCE_FAITHFUL", "LISTENING_ADAPTATION"]
    model_choice: Literal["AUTO", "ASTRA_MEDIUM", "ASTRA_HIGH", "ASTRA_XHIGH", "SOL"] = "AUTO"
    max_cost_usd_per_request: float = Field(default=1.0, gt=0, le=20)
    max_total_cost_usd: float = Field(default=10.0, gt=0, le=500)
    max_requests: int = Field(default=20, ge=1, le=200)
    reading_docx: bool = True
    litres_docx: bool = False
    pronunciation_dictionary: bool = True
    owner_authorizes_paid_requests: Literal[True]


class ExistingAudioApprovalRequest(BaseModel):
    human_actor: str = Field(min_length=1, max_length=300)
    accepted_attention_codes: list[str] = Field(default_factory=list, max_length=100)
    reading_docx: bool = True
    litres_docx: bool = False
    pronunciation_dictionary: bool = True


class AudioScriptRevisionRequest(BaseModel):
    content: AudioScriptContent
    human_actor: str = Field(min_length=1, max_length=300)
    change_summary: str = Field(min_length=3, max_length=4000)


class ExistingAudioPrepareView(BaseModel):
    run_id: str
    audio_script: AudioScriptView
    artifacts: list[dict[str, Any]] = Field(default_factory=list)


def _extract_source(filename: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.casefold()
    try:
        if suffix in {".txt", ".md"}:
            value = payload.decode("utf-8")
        elif suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(payload))
            paragraphs = [item.text for item in document.paragraphs if item.text.strip()]
            for table in document.tables:
                paragraphs.append("Таблица:")
                paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            value = "\n\n".join(paragraphs)
        elif suffix == ".pdf":
            from pypdf import PdfReader

            pages = [page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages]
            if not pages or any(not page.strip() for page in pages):
                raise AudioScriptGateError(
                    "PDF extraction is partial; audio adaptation cannot silently omit a page"
                )
            value = "\n\n".join(pages)
        elif suffix == ".rtf":
            raw = payload.decode("utf-8", errors="replace")
            value = " ".join(
                part for part in raw.replace("\\par", "\n").split() if not part.startswith("\\")
            )
        elif suffix == ".epub":
            with ZipFile(BytesIO(payload)) as archive:
                html_parts = [
                    archive.read(name).decode("utf-8", errors="replace")
                    for name in archive.namelist()
                    if name.casefold().endswith((".xhtml", ".html", ".htm"))
                ]
            value = unescape(
                "\n\n".join(
                    re.sub(r"<[^>]+>", " ", part).replace("&nbsp;", " ") for part in html_parts
                )
            )
        else:
            raise AudioScriptGateError(
                "supported source types: UTF-8 TXT, MD, DOCX, PDF, RTF, EPUB"
            )
    except (UnicodeDecodeError, BadZipFile, OSError, PdfReadError, ValueError) as exc:
        raise AudioScriptGateError(
            "source extraction failed completely and cannot be guessed"
        ) from exc
    normalized = "\n\n".join(part.strip() for part in value.splitlines() if part.strip())
    if len(normalized) < 20:
        raise AudioScriptGateError("source extraction produced too little text")
    return normalized


def _chunks(value: str, *, maximum: int = 18_000) -> list[str]:
    paragraphs = [item.strip() for item in value.split("\n\n") if item.strip()]
    chunks: list[str] = []
    current: list[str] = []
    length = 0
    for paragraph in paragraphs:
        if current and length + len(paragraph) > maximum:
            chunks.append("\n\n".join(current))
            current = []
            length = 0
        current.append(paragraph)
        length += len(paragraph) + 2
    if current:
        chunks.append("\n\n".join(current))
    return chunks


def build_audio_script_router(
    data_dir: Path,
    require_token: Callable[..., None],
    gateway: ModelGateway,
) -> APIRouter:
    scripts = AudioScriptService(data_dir)
    runtime = DurableAutoBookRuntime(data_dir)
    routing = ModelRoutingService(data_dir)
    exporter = AutoBookExporter(data_dir, runtime)
    router = APIRouter(dependencies=[Depends(require_token)])

    def _choice(book_id: str, value: str) -> tuple[str, str | None, str]:
        if value == "AUTO":
            choice = routing.resolve(
                book_id,
                "SECTION_DRAFT",
                provider="openai",
                selection_mode="AUTO",
                selection_scope=None,
                model=None,
                quality_risk="HIGH",
            )
            return choice.model, choice.reasoning_effort, choice.rationale
        mapping = {
            "ASTRA_MEDIUM": ("gpt-6-astra", "medium"),
            "ASTRA_HIGH": ("gpt-6-astra", "high"),
            "ASTRA_XHIGH": ("gpt-6-astra", "xhigh"),
            "SOL": ("gpt-5.6-sol", None),
        }
        model, effort = mapping[value]
        return model, effort, "Human manual model choice for the audio-edit operation"

    @router.get("/api/projects/{book_id}/audio-scripts")
    def list_audio_scripts(book_id: str) -> list[dict[str, Any]]:
        try:
            return [item.model_dump(mode="json") for item in scripts.list_scripts(book_id)]
        except AudioScriptError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @router.post("/api/projects/{book_id}/audio-scripts/prepare")
    def prepare_existing_audio(
        book_id: str,
        payload: ExistingAudioPrepareRequest,
    ) -> dict[str, Any]:
        if payload.max_total_cost_usd < payload.max_cost_usd_per_request:
            raise HTTPException(status_code=409, detail="total budget must cover one request")
        try:
            raw = base64.b64decode(payload.source_content_base64, validate=True)
        except (ValueError, binascii.Error) as exc:
            raise HTTPException(status_code=400, detail="source is not valid base64") from exc
        if not raw or len(raw) > 25_000_000:
            raise HTTPException(status_code=400, detail="source must be between 1 byte and 25 MB")
        worker_id: str | None = None
        run_id: str | None = None
        try:
            source_text = _extract_source(payload.source_filename, raw)
            source_hash = hashlib.sha256(raw).hexdigest()
            safe_suffix = Path(payload.source_filename).suffix.casefold()
            source_dir = runtime.projects.projects_dir / book_id / "audio-sources"
            source_dir.mkdir(parents=True, exist_ok=True)
            source_path = source_dir / f"{source_hash}{safe_suffix}"
            if not source_path.exists():
                source_path.write_bytes(raw)
            source_relative_path = str(
                source_path.relative_to(runtime.projects.projects_dir / book_id)
            )
            chapters = [
                MasterChapter(
                    chapter_id=f"imported-section-{index}",
                    title=f"Раздел {index}",
                    paragraphs=[chunk],
                )
                for index, chunk in enumerate(_chunks(source_text), start=1)
            ]
            if len(chapters) > payload.max_requests:
                raise AudioScriptGateError(
                    "request limit is lower than the number of bounded source sections"
                )
            master = StructuredBookMaster(
                title=payload.title,
                author=payload.author,
                chapters=chapters,
            )
            existing = next(
                (
                    item
                    for item in scripts.list_scripts(book_id)
                    if item.source_hash == source_hash
                    and item.adaptation_mode == payload.adaptation_mode
                    and item.provenance.get("workflow") == "EXISTING_TEXT_TO_AUDIO"
                    and item.content.title == payload.title
                    and item.content.author == payload.author
                    and not item.stale_against_source
                ),
                None,
            )
            if existing is not None:
                run_id = str(existing.provenance["run_id"])
                existing_artifacts: list[dict[str, Any]] = []
                if existing.ready_for_export:
                    source_shaped = StructuredBookMaster(
                        title=existing.content.title,
                        author=existing.content.author,
                        language=existing.content.language,
                        chapters=[
                            MasterChapter(
                                chapter_id=section.source_chapter_id,
                                title=section.title,
                                paragraphs=section.paragraphs,
                            )
                            for section in existing.content.sections
                        ],
                    )
                    selection = AutoBookOutputSelection(
                        full_manuscript_docx=False,
                        audio_reading_docx=payload.reading_docx,
                        audio_litres_docx=payload.litres_docx,
                        voice_text_txt=True,
                        pronunciation_dictionary=payload.pronunciation_dictionary,
                    )
                    bundle = exporter.export_selected(
                        book_id,
                        run_id,
                        source_shaped,
                        selection,
                        audio_script=existing,
                    )
                    existing_artifacts = [item.model_dump(mode="json") for item in bundle.artifacts]
                return ExistingAudioPrepareView(
                    run_id=run_id,
                    audio_script=existing,
                    artifacts=existing_artifacts,
                ).model_dump(mode="json")

            intent = AutoBookIntent(
                idea=f"Prepare existing source {payload.source_filename} for audio",
                reader_hint=f"audio_adaptation_mode={payload.adaptation_mode}",
                author_name=payload.author,
                outputs=AutoBookOutputSelection(
                    full_manuscript_docx=False,
                    audio_reading_docx=payload.reading_docx,
                    audio_litres_docx=payload.litres_docx,
                    voice_text_txt=True,
                    pronunciation_dictionary=payload.pronunciation_dictionary,
                ),
                attachments=[
                    AutoBookAttachment(
                        path=source_relative_path,
                        role="SOURCE",
                        content_hash=source_hash,
                    )
                ],
                max_cost_usd_per_request=payload.max_cost_usd_per_request,
                max_total_cost_usd=payload.max_total_cost_usd,
                max_requests=payload.max_requests,
            )
            latest = runtime.latest(book_id)
            matching_run = (
                latest
                if latest is not None
                and any(
                    item.role == "SOURCE"
                    and item.content_hash == source_hash
                    and item.path == source_relative_path
                    for item in latest.intent.attachments
                )
                and latest.intent.reader_hint == f"audio_adaptation_mode={payload.adaptation_mode}"
                else None
            )
            if matching_run is not None and matching_run.status == "UNKNOWN_OUTCOME":
                raise HTTPException(
                    status_code=503,
                    detail=("A previous provider outcome is unknown; automatic retry is blocked."),
                )
            run = matching_run or runtime.create_run(
                book_id,
                intent,
                authorized_by=payload.author,
            )
            run_id = run.run_id
            worker_id = f"audio-script:{run.run_id}"
            runtime.claim(book_id, run.run_id, worker_id)
            runtime.set_stage(
                book_id,
                run.run_id,
                AutoBookStage.AUDIO_EDITORIAL,
                message="Готовая книга перерабатывается в отдельный AudioScript",
            )
            adapted: dict[str, str] = {}
            model_runs: list[dict[str, Any]] = []
            for ordinal, chapter in enumerate(master.chapters, start=1):
                model, effort, rationale = _choice(book_id, payload.model_choice)
                input_payload = {
                    "source_hash": source_hash,
                    "source_chapter_id": chapter.chapter_id,
                    "source_text": chapter.paragraphs,
                    "adaptation_mode": payload.adaptation_mode,
                }
                operation = runtime.ensure_operation(
                    book_id,
                    run.run_id,
                    ordinal=ordinal,
                    stage=AutoBookStage.AUDIO_EDITORIAL,
                    operation=f"audio-edit:{chapter.chapter_id}",
                    input_payload=input_payload,
                    provider="openai",
                    model=model,
                    reasoning_effort=effort,
                    estimated_cost_usd=payload.max_cost_usd_per_request,
                )
                if operation.state == "SUCCEEDED" and operation.output is not None:
                    adapted[chapter.chapter_id] = str(operation.output["text"])
                    model_runs.append(
                        {
                            "provider": operation.provider,
                            "model": operation.model,
                            "reasoning_effort": operation.reasoning_effort,
                            "provider_run_id": operation.provider_run_id,
                            "usage": operation.output.get("usage", {}),
                            "source_chapter_id": chapter.chapter_id,
                            "prompt_hash": AUDIO_SCRIPT_EDITOR_V1.prompt_hash,
                            "recovered_idempotently": True,
                        }
                    )
                    continue
                if operation.state == "UNKNOWN":
                    raise HTTPException(
                        status_code=503,
                        detail=(
                            "A previous provider outcome is unknown; automatic retry is blocked."
                        ),
                    )
                if operation.state in {"RESERVED", "RUNNING"}:
                    runtime.mark_unknown(
                        book_id,
                        run.run_id,
                        operation.operation_id,
                        provider_run_id=operation.provider_run_id,
                    )
                    raise HTTPException(
                        status_code=503,
                        detail=(
                            "An interrupted provider operation has an unknown outcome; "
                            "automatic retry is blocked."
                        ),
                    )
                runtime.reserve(
                    book_id,
                    run.run_id,
                    operation.operation_id,
                    payload.max_cost_usd_per_request,
                )
                try:
                    result = gateway.generate(
                        ModelTaskRequest(
                            task_id=operation.operation_id,
                            task_type="SECTION_DRAFT",
                            role="WRITER",
                            provider="openai",
                            model=model,
                            prompt_id=AUDIO_SCRIPT_EDITOR_V1.prompt_id,
                            prompt_version=AUDIO_SCRIPT_EDITOR_V1.version,
                            prompt_hash=AUDIO_SCRIPT_EDITOR_V1.prompt_hash,
                            section_objective=(
                                "Prepare this exact existing-book section for professional "
                                f"listening in {payload.adaptation_mode} mode."
                            ),
                            authoritative_context=input_payload,
                            task_payload={
                                "workflow": "EXISTING_TEXT_TO_AUDIO",
                                "adaptation_mode": payload.adaptation_mode,
                                "selection_mode": (
                                    "AUTO" if payload.model_choice == "AUTO" else "MANUAL"
                                ),
                                "routing_rationale": rationale,
                            },
                            reasoning_effort=cast(Any, effort),
                            max_output_tokens=12_000,
                            max_cost_usd=payload.max_cost_usd_per_request,
                        ),
                        AUDIO_SCRIPT_EDITOR_V1,
                    )
                    output = SectionDraftOutput.model_validate(result.output)
                except (
                    httpx.TransportError,
                    ModelOutputError,
                    ModelProviderError,
                    ValueError,
                ) as exc:
                    runtime.mark_unknown(
                        book_id,
                        run.run_id,
                        operation.operation_id,
                        provider_run_id=None,
                    )
                    raise HTTPException(
                        status_code=503,
                        detail=(
                            "Provider outcome is unknown; the request is not retried blindly and "
                            "its reservation remains accounted for."
                        ),
                    ) from exc
                confirmed = min(
                    payload.max_cost_usd_per_request,
                    max(0.0, float(result.usage.get("cost_usd", 0.0))),
                )
                runtime.complete_operation(
                    book_id,
                    run.run_id,
                    operation.operation_id,
                    output={
                        "text": output.text,
                        "notes": output.notes,
                        "usage": result.usage,
                    },
                    confirmed_cost_usd=confirmed,
                    provider_run_id=result.provider_run_id,
                )
                adapted[chapter.chapter_id] = output.text
                model_runs.append(
                    {
                        "provider": "openai",
                        "model": model,
                        "reasoning_effort": effort,
                        "provider_run_id": result.provider_run_id,
                        "usage": result.usage,
                        "source_chapter_id": chapter.chapter_id,
                        "prompt_hash": AUDIO_SCRIPT_EDITOR_V1.prompt_hash,
                    }
                )
            content, transformations = scripts.content_from_master(
                master,
                adaptation_mode=payload.adaptation_mode,
                adapted_chapters=adapted,
            )
            proposed = scripts.create_proposal(
                book_id,
                source_kind="IMPORTED_SOURCE",
                source_identity=source_relative_path,
                source_hash=source_hash,
                adaptation_mode=payload.adaptation_mode,
                content=content,
                transformations=transformations,
                provenance={
                    "workflow": "EXISTING_TEXT_TO_AUDIO",
                    "run_id": run.run_id,
                    "source_filename": Path(payload.source_filename).name,
                    "source_hash": source_hash,
                    "model_runs": model_runs,
                },
            )
            runtime.complete_stage(
                book_id,
                run.run_id,
                AutoBookStage.AUDIO_EDITORIAL,
                evidence={
                    "audio_script_id": proposed.audio_script_id,
                    "source_hash": source_hash,
                    "authority_status": proposed.status,
                    "human_approval_required": True,
                },
                message="AudioScript предложен и ждёт решения человека",
            )
            runtime.set_stage(
                book_id,
                run.run_id,
                AutoBookStage.AUDIO_EDITORIAL,
                status="MANUSCRIPT_READY",
                message="AudioScript предложен и ждёт решения человека",
            )
            runtime.release(book_id, run.run_id, worker_id)
            return ExistingAudioPrepareView(
                run_id=run.run_id,
                audio_script=proposed,
            ).model_dump(mode="json")
        except HTTPException:
            if worker_id is not None and run_id is not None:
                runtime.release(book_id, run_id, worker_id)
            raise
        except (
            AudioScriptError,
            AutoBookRuntimeError,
            ModelOutputError,
            ModelProviderError,
        ) as exc:
            if worker_id is not None and run_id is not None:
                runtime.release(book_id, run_id, worker_id)
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @router.post("/api/projects/{book_id}/audio-scripts/{audio_script_id}/approve")
    def approve_existing_audio(
        book_id: str,
        audio_script_id: str,
        payload: ExistingAudioApprovalRequest,
    ) -> dict[str, Any]:
        try:
            script = scripts.approve(
                book_id,
                audio_script_id,
                human_actor=payload.human_actor,
                accepted_attention_codes=payload.accepted_attention_codes,
            )
            run_id = str(script.provenance.get("run_id", ""))
            if not run_id:
                raise AudioScriptGateError("audio workflow run identity is missing")
            selection = AutoBookOutputSelection(
                full_manuscript_docx=False,
                audio_reading_docx=payload.reading_docx,
                audio_litres_docx=payload.litres_docx,
                voice_text_txt=True,
                pronunciation_dictionary=payload.pronunciation_dictionary,
            )
            source_shaped = StructuredBookMaster(
                title=script.content.title,
                author=script.content.author,
                language=script.content.language,
                chapters=[
                    MasterChapter(
                        chapter_id=section.source_chapter_id,
                        title=section.title,
                        paragraphs=section.paragraphs,
                    )
                    for section in script.content.sections
                ],
            )
            bundle = exporter.export_selected(
                book_id,
                run_id,
                source_shaped,
                selection,
                audio_script=script,
            )
            runtime.set_stage(
                book_id,
                run_id,
                AutoBookStage.MASTER_AND_EXPORTS,
                status="PACKAGE_READY",
                message="Утверждённый AudioScript и production handoff готовы",
            )
            return {
                "audio_script": script.model_dump(mode="json"),
                "artifacts": [item.model_dump(mode="json") for item in bundle.artifacts],
                "output_directory": bundle.output_directory,
            }
        except (AudioScriptError, AutoBookRuntimeError, AutoBookBudgetError) as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    @router.put("/api/projects/{book_id}/audio-scripts/{audio_script_id}")
    def revise_existing_audio(
        book_id: str,
        audio_script_id: str,
        payload: AudioScriptRevisionRequest,
    ) -> dict[str, Any]:
        try:
            return scripts.revise_by_human(
                book_id,
                audio_script_id,
                content=payload.content,
                human_actor=payload.human_actor,
                change_summary=payload.change_summary,
            ).model_dump(mode="json")
        except AudioScriptError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc

    return router
