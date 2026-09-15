from __future__ import annotations

import hashlib
from html import escape
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from typing import Any, Literal

from docx import Document
from docx.document import Document as DocumentObject
from docx.opc.constants import CONTENT_TYPE as DOCX_CONTENT_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as DOCX_RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
import ebooklib  # type: ignore[import-untyped]
from ebooklib import epub
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from pydantic import BaseModel, Field, model_validator
from reportlab.lib.enums import TA_CENTER  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore[import-untyped]
from reportlab.lib.units import mm  # type: ignore[import-untyped]
from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
from reportlab.pdfbase.ttfonts import TTFont  # type: ignore[import-untyped]
from reportlab.platypus import Image as PdfImage  # type: ignore[import-untyped]
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib import colors  # type: ignore[import-untyped]

from .authority import canonical_json
from .audio_script import (
    AudioScriptContent,
    AudioScriptGateError,
    AudioScriptSection,
    AudioScriptService,
    AudioScriptView,
)
from .auto_book_runtime import (
    AutoBookArtifactView,
    AutoBookOutputKind,
    AutoBookOutputSelection,
    DurableAutoBookRuntime,
)


class AutoBookExportError(RuntimeError):
    pass


class MasterTable(BaseModel):
    object_id: str = Field(min_length=1, max_length=160)
    title: str = Field(min_length=1, max_length=500)
    headers: list[str] = Field(min_length=1, max_length=30)
    rows: list[list[str]] = Field(min_length=1, max_length=2000)
    audio_equivalent: str = Field(min_length=1, max_length=12000)
    source_note: str | None = Field(default=None, max_length=2000)
    placement_after_paragraph: int | None = Field(default=None, ge=0)

    @model_validator(mode="after")
    def rectangular(self) -> MasterTable:
        if any(len(row) != len(self.headers) for row in self.rows):
            raise ValueError("every table row must match the header width")
        return self


class MasterVisual(BaseModel):
    object_id: str = Field(min_length=1, max_length=160)
    kind: Literal["CHART", "SCHEME", "ILLUSTRATION"]
    title: str = Field(min_length=1, max_length=500)
    caption: str = Field(min_length=1, max_length=2000)
    alt_text: str = Field(min_length=1, max_length=4000)
    audio_equivalent: str = Field(min_length=1, max_length=12000)
    data: list[tuple[str, float]] = Field(default_factory=list, max_length=100)
    source_note: str | None = Field(default=None, max_length=2000)
    rights_note: str = Field(default="Created programmatically by BOOK OS", max_length=1000)
    placement_after_paragraph: int | None = Field(default=None, ge=0)


class MasterChapter(BaseModel):
    chapter_id: str = Field(min_length=1, max_length=160)
    title: str = Field(min_length=1, max_length=500)
    paragraphs: list[str] = Field(min_length=1)
    tables: list[MasterTable] = Field(default_factory=list)
    visuals: list[MasterVisual] = Field(default_factory=list)


class StructuredBookMaster(BaseModel):
    title: str = Field(min_length=1, max_length=500)
    author: str = Field(min_length=1, max_length=300)
    language: str = "ru"
    chapters: list[MasterChapter] = Field(min_length=1)
    bibliography: list[str] = Field(default_factory=list)
    publisher_annotation: str = Field(default="", max_length=12000)

    @property
    def manifest_hash(self) -> str:
        import hashlib

        return hashlib.sha256(
            canonical_json(self.model_dump(mode="json")).encode("utf-8")
        ).hexdigest()


class ExportBundle(BaseModel):
    master_hash: str
    artifacts: list[AutoBookArtifactView]
    output_directory: str


class AutoBookExporter:
    EXPORTER_VERSION = "auto-book-export.v1.1.0"
    LITRES_PROFILE_VERSION = "litres-ebook.2026-09-15.v2"
    LITRES_PROFILE_CHECKED_AT = "2026-09-15"
    PROFILE_VERSIONS: dict[AutoBookOutputKind, str] = {
        "FULL_MANUSCRIPT_DOCX": "full-manuscript.v1",
        "LITRES_EBOOK_DOCX": "litres-ebook.2026-09-15.v2",
        "READING_PDF": "reading-pdf.v1",
        "EPUB": "epub3.v1",
        "AUDIO_READING_DOCX": "audio-reading.v1",
        "AUDIO_LITRES_DOCX": "audio-litres.owner-profile.v1",
        "VOICE_TEXT_TXT": "voice-text.v1",
        "PRONUNCIATION_DICTIONARY": "pronunciation.v1",
        "AUDIO_PRODUCTION_HANDOFF": "book-os-audiobook-handoff.v2",
        "READER_EXTRAS": "reader-extras.v1",
        "PUBLISHER_PACK": "publisher-pack.v1",
    }

    _FILE_NAMES: dict[AutoBookOutputKind, str] = {
        "FULL_MANUSCRIPT_DOCX": "Полная-рукопись.docx",
        "LITRES_EBOOK_DOCX": "Электронная-версия-Литрес.docx",
        "READING_PDF": "Версия-для-чтения.pdf",
        "EPUB": "Электронная-книга.epub",
        "AUDIO_READING_DOCX": "Аудиоредакция-для-чтения.docx",
        "AUDIO_LITRES_DOCX": "Аудиоредакция-для-Литрес.docx",
        "VOICE_TEXT_TXT": "Текст-для-озвучки.txt",
        "PRONUNCIATION_DICTIONARY": "Словарь-произношения.txt",
        "AUDIO_PRODUCTION_HANDOFF": "Audiobook-Studio-handoff.json",
        "READER_EXTRAS": "Дополнительные-материалы.docx",
        "PUBLISHER_PACK": "Издательский-пакет.json",
    }

    def __init__(self, data_dir: Path, runtime: DurableAutoBookRuntime) -> None:
        self.data_dir = data_dir
        self.runtime = runtime

    @staticmethod
    def _font_path() -> Path | None:
        candidates = (
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        )
        return next((path for path in candidates if path.is_file()), None)

    @classmethod
    def _register_pdf_font(cls) -> str:
        path = cls._font_path()
        if path is None:
            raise AutoBookExportError("a Unicode TrueType font is required for Russian PDF export")
        name = "BookOSUnicode"
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
        return name

    @classmethod
    def _image_font(cls, size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
        path = cls._font_path()
        if path is None:
            return ImageFont.load_default()
        return ImageFont.truetype(str(path), size=size)

    @staticmethod
    def _apply_docx_styles(document: DocumentObject) -> None:
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)
        for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
            style = document.styles[style_name]
            style.font.name = "Arial"
            style.font.size = Pt(size)

    @staticmethod
    def _add_bibliography(document: DocumentObject, bibliography: list[str]) -> None:
        if not bibliography:
            return
        document.add_page_break()  # type: ignore[no-untyped-call]
        document.add_heading("Библиография", level=1)
        for index, entry in enumerate(bibliography, start=1):
            document.add_paragraph(f"{index}. {entry}")

    def _render_visual(self, visual: MasterVisual, output_dir: Path) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        output = output_dir / f"{visual.object_id}.png"
        image = Image.new("RGB", (1024, 576), "white")
        draw = ImageDraw.Draw(image)
        title_font = self._image_font(42)
        body_font = self._image_font(28)
        draw.text((80, 55), visual.title, fill="#17253f", font=title_font)
        if visual.kind == "CHART" and visual.data:
            maximum = max(value for _, value in visual.data) or 1.0
            bar_width = max(44, 760 // len(visual.data))
            for index, (label, value) in enumerate(visual.data):
                x0 = 100 + index * bar_width
                height = int(330 * value / maximum)
                y0 = 455 - height
                draw.rectangle((x0, y0, x0 + bar_width - 18, 455), fill="#2d6cdf")
                draw.text((x0, 468), label, fill="#17253f", font=body_font)
                draw.text((x0, max(115, y0 - 40)), str(value), fill="#17253f", font=body_font)
        else:
            draw.rounded_rectangle((90, 145, 934, 430), radius=24, outline="#2d6cdf", width=6)
            draw.multiline_text(
                (130, 205),
                visual.alt_text,
                fill="#17253f",
                font=body_font,
                spacing=18,
            )
        image.save(output, format="PNG", optimize=True, dpi=(72, 72))
        return output

    @staticmethod
    def _reader_source_note(value: str | None) -> str | None:
        if value is None:
            return None
        cleaned = value.strip()
        if not cleaned or cleaned.startswith("revision:"):
            return None
        return cleaned

    @staticmethod
    def _plain_inline(value: str) -> str:
        value = re.sub(r"\[([^\]]+)\]\(#([^)]+)\)", r"\1", value)
        value = re.sub(r"\[\^([^\]]+)\]", r"\1", value)
        value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
        value = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", value)
        return value

    @staticmethod
    def _heading_parts(value: str) -> tuple[int, str, str | None] | None:
        match = re.fullmatch(r"(#{2,4})\s+(.+?)(?:\s+\{#([A-Za-z0-9_-]+)\})?", value.strip())
        if match is None:
            return None
        return len(match.group(1)), match.group(2).strip(), match.group(3)

    @classmethod
    def _semantic_kind(cls, value: str) -> tuple[str, Any]:
        raw = value.strip()
        heading = cls._heading_parts(raw)
        if heading is not None:
            return "HEADING", heading
        footnote = re.fullmatch(r"\[\^([^\]]+)\]:\s*(.+)", raw, flags=re.DOTALL)
        if footnote is not None:
            return "FOOTNOTE", (footnote.group(1), footnote.group(2).strip())
        lines = [line.strip() for line in raw.splitlines() if line.strip()]
        if lines and all(re.match(r"^[-*]\s+", line) for line in lines):
            return "BULLET_LIST", [re.sub(r"^[-*]\s+", "", line) for line in lines]
        if lines and all(re.match(r"^\d+[.)]\s+", line) for line in lines):
            return "NUMBERED_LIST", [re.sub(r"^\d+[.)]\s+", "", line) for line in lines]
        if raw.startswith(">"):
            return "NOTE", " ".join(line.lstrip("> ") for line in raw.splitlines())
        return "PARAGRAPH", raw

    @classmethod
    def _semantic_fragments(cls, value: str) -> list[str]:
        kind, payload = cls._semantic_kind(value)
        if kind == "HEADING":
            return [cls._plain_inline(payload[1])]
        if kind in {"BULLET_LIST", "NUMBERED_LIST"}:
            return [cls._plain_inline(item) for item in payload]
        if kind == "FOOTNOTE":
            return [cls._plain_inline(payload[1])]
        return [cls._plain_inline(str(payload))]

    @classmethod
    def _footnote_definitions(cls, master: StructuredBookMaster) -> dict[str, str]:
        definitions: dict[str, str] = {}
        for chapter in master.chapters:
            for paragraph in chapter.paragraphs:
                kind, payload = cls._semantic_kind(paragraph)
                if kind != "FOOTNOTE":
                    continue
                note_id, text = payload
                if note_id in definitions:
                    raise AutoBookExportError(f"duplicate footnote definition: {note_id}")
                definitions[note_id] = text
        return definitions

    @classmethod
    def _validate_footnote_references(
        cls, master: StructuredBookMaster, definitions: dict[str, str]
    ) -> None:
        referenced: set[str] = set()
        for chapter in master.chapters:
            for paragraph in chapter.paragraphs:
                kind, _payload = cls._semantic_kind(paragraph)
                if kind == "FOOTNOTE":
                    continue
                referenced.update(re.findall(r"\[\^([^\]]+)\]", paragraph))
        missing = sorted(referenced - set(definitions))
        if missing:
            raise AutoBookExportError(
                "footnote references have no definitions: " + ", ".join(missing)
            )

    @staticmethod
    def _docx_add_bookmark(paragraph: Any, anchor: str, bookmark_id: int) -> None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), anchor)
        paragraph._p.append(start)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        paragraph._p.append(end)

    @staticmethod
    def _docx_add_hyperlink(paragraph: Any, label: str, anchor: str) -> None:
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), anchor)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.extend([color, underline])
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    @staticmethod
    def _docx_add_footnote_reference(paragraph: Any, footnote_id: int) -> None:
        run = OxmlElement("w:r")
        reference = OxmlElement("w:footnoteReference")
        reference.set(qn("w:id"), str(footnote_id))
        run.append(reference)
        paragraph._p.append(run)

    @classmethod
    def _docx_add_inline(cls, paragraph: Any, value: str, footnote_ids: dict[str, int]) -> None:
        token = re.compile(
            r"(\[[^\]]+\]\(#[^)]+\)|\[\^[^\]]+\]|\*\*[^*]+\*\*|(?<!\*)\*[^*\n]+\*(?!\*))"
        )
        cursor = 0
        for match in token.finditer(value):
            if match.start() > cursor:
                paragraph.add_run(value[cursor : match.start()])
            marked = match.group(0)
            link = re.fullmatch(r"\[([^\]]+)\]\(#([^)]+)\)", marked)
            footnote = re.fullmatch(r"\[\^([^\]]+)\]", marked)
            if link is not None:
                cls._docx_add_hyperlink(paragraph, link.group(1), link.group(2))
            elif footnote is not None:
                note_id = footnote.group(1)
                if note_id not in footnote_ids:
                    raise AutoBookExportError(f"undefined footnote reference: {note_id}")
                cls._docx_add_footnote_reference(paragraph, footnote_ids[note_id])
            elif marked.startswith("**"):
                run = paragraph.add_run(marked[2:-2])
                run.bold = True
            else:
                run = paragraph.add_run(marked[1:-1])
                run.italic = True
            cursor = match.end()
        if cursor < len(value):
            paragraph.add_run(value[cursor:])

    @classmethod
    def _docx_add_semantic(
        cls,
        document: DocumentObject,
        value: str,
        footnote_ids: dict[str, int],
        bookmark_ids: dict[str, int],
    ) -> None:
        kind, payload = cls._semantic_kind(value)
        if kind == "FOOTNOTE":
            return
        if kind == "HEADING":
            level, text, anchor = payload
            paragraph = document.add_heading(level=min(max(level, 2), 3))
            if anchor:
                bookmark_id = bookmark_ids.setdefault(anchor, len(bookmark_ids) + 1)
                cls._docx_add_bookmark(paragraph, anchor, bookmark_id)
            cls._docx_add_inline(paragraph, text, footnote_ids)
            return
        if kind in {"BULLET_LIST", "NUMBERED_LIST"}:
            style = "List Bullet" if kind == "BULLET_LIST" else "List Number"
            for item in payload:
                paragraph = document.add_paragraph(style=style)
                cls._docx_add_inline(paragraph, item, footnote_ids)
            return
        if kind == "NOTE":
            paragraph = document.add_paragraph(style="Quote")
            cls._docx_add_inline(paragraph, str(payload), footnote_ids)
            return
        paragraph = document.add_paragraph()
        cls._docx_add_inline(paragraph, str(payload), footnote_ids)

    @classmethod
    def _attach_docx_footnotes(
        cls,
        document: DocumentObject,
        definitions: dict[str, str],
        footnote_ids: dict[str, int],
    ) -> None:
        if not definitions:
            return
        rows = [
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>',
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>',
        ]
        for note_key, text in definitions.items():
            note_id = footnote_ids[note_key]
            clean = escape(cls._plain_inline(text))
            rows.append(
                f'<w:footnote w:id="{note_id}"><w:p><w:r><w:footnoteRef/></w:r>'
                f'<w:r><w:t xml:space="preserve"> {clean}</w:t></w:r></w:p></w:footnote>'
            )
        xml = f"<w:footnotes {nsdecls('w')}>{''.join(rows)}</w:footnotes>"
        part = XmlPart(
            PackURI("/word/footnotes.xml"),
            DOCX_CONTENT_TYPE.WML_FOOTNOTES,
            parse_xml(xml),
            document.part.package,
        )
        document.part.relate_to(part, DOCX_RELATIONSHIP_TYPE.FOOTNOTES)

    @classmethod
    def _html_inline(cls, value: str) -> str:
        rendered = escape(value)
        rendered = re.sub(
            r"\[([^\]]+)\]\(#([A-Za-z0-9_-]+)\)",
            r'<a href="#\2">\1</a>',
            rendered,
        )
        rendered = re.sub(
            r"\[\^([^\]]+)\]",
            r'<sup id="fnref-\1"><a href="#fn-\1">\1</a></sup>',
            rendered,
        )
        rendered = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", rendered)
        rendered = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"<em>\1</em>", rendered)
        return rendered

    @classmethod
    def _epub_semantic_html(cls, value: str) -> str:
        kind, payload = cls._semantic_kind(value)
        if kind == "HEADING":
            level, text, anchor = payload
            anchor_attr = f' id="{escape(anchor)}"' if anchor else ""
            return f"<h{level}{anchor_attr}>{cls._html_inline(text)}</h{level}>"
        if kind in {"BULLET_LIST", "NUMBERED_LIST"}:
            tag = "ul" if kind == "BULLET_LIST" else "ol"
            return (
                f"<{tag}>"
                + "".join(f"<li>{cls._html_inline(item)}</li>" for item in payload)
                + f"</{tag}>"
            )
        if kind == "NOTE":
            return f'<aside class="note"><p>{cls._html_inline(str(payload))}</p></aside>'
        if kind == "FOOTNOTE":
            note_id, text = payload
            return (
                f'<aside class="footnote" id="fn-{escape(note_id)}">'
                f"<p><sup>{escape(note_id)}</sup> {cls._html_inline(text)} "
                f'<a href="#fnref-{escape(note_id)}" aria-label="Назад к тексту">↩</a>'
                "</p></aside>"
            )
        return f"<p>{cls._html_inline(str(payload))}</p>"

    @classmethod
    def _pdf_inline_html(cls, value: str) -> str:
        rendered = escape(value)
        rendered = re.sub(
            r"\[([^\]]+)\]\(#([A-Za-z0-9_-]+)\)",
            r'<a href="#\2" color="#0563C1">\1</a>',
            rendered,
        )
        rendered = re.sub(r"\[\^([^\]]+)\]", r"[\1]", rendered)
        rendered = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", rendered)
        rendered = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"<i>\1</i>", rendered)
        return rendered

    @classmethod
    def _pdf_semantic_flowables(
        cls, value: str, body: ParagraphStyle, heading: ParagraphStyle
    ) -> list[Any]:
        kind, payload = cls._semantic_kind(value)
        if kind == "HEADING":
            _level, text, anchor = payload
            anchor_markup = f'<a name="{escape(anchor)}"/>' if anchor else ""
            return [
                Paragraph(f"{anchor_markup}<b>{cls._pdf_inline_html(text)}</b>", heading),
                Spacer(1, 2 * mm),
            ]
        if kind in {"BULLET_LIST", "NUMBERED_LIST"}:
            output: list[Any] = []
            for index, item in enumerate(payload, start=1):
                prefix = "•" if kind == "BULLET_LIST" else f"{index}."
                output.extend(
                    [
                        Paragraph(f"{escape(prefix)} {cls._pdf_inline_html(item)}", body),
                        Spacer(1, 1.5 * mm),
                    ]
                )
            return output
        if kind == "NOTE":
            return [
                Paragraph(f"<i>Примечание.</i> {cls._pdf_inline_html(str(payload))}", body),
                Spacer(1, 2 * mm),
            ]
        if kind == "FOOTNOTE":
            note_id, text = payload
            return [
                Paragraph(f"[{escape(note_id)}] {cls._pdf_inline_html(text)}", body),
                Spacer(1, 2 * mm),
            ]
        return [Paragraph(cls._pdf_inline_html(str(payload)), body), Spacer(1, 2.5 * mm)]

    @staticmethod
    def _ordered_chapter_blocks(chapter: MasterChapter) -> list[tuple[str, Any]]:
        slots: dict[int, list[tuple[str, Any]]] = {}
        tail: list[tuple[str, Any]] = []
        paragraph_count = len(chapter.paragraphs)
        for kind, values in (("TABLE", chapter.tables), ("VISUAL", chapter.visuals)):
            for value in values:
                placement = value.placement_after_paragraph
                if placement is None or placement > paragraph_count:
                    tail.append((kind, value))
                else:
                    slots.setdefault(max(0, placement), []).append((kind, value))
        blocks: list[tuple[str, Any]] = [*slots.get(0, [])]
        for index, paragraph in enumerate(chapter.paragraphs, start=1):
            blocks.append(("PARAGRAPH", paragraph))
            blocks.extend(slots.get(index, []))
        blocks.extend(tail)
        return blocks

    def _docx(
        self,
        master: StructuredBookMaster,
        output: Path,
        *,
        audio: bool,
        include_extras_only: bool = False,
        visual_dir: Path,
    ) -> dict[str, Any]:
        document = Document()
        self._apply_docx_styles(document)
        document.add_heading(master.title, level=0)
        document.add_paragraph(master.author)
        expected_tables = 0
        expected_visuals = 0
        footnote_definitions = {} if include_extras_only else self._footnote_definitions(master)
        if not include_extras_only:
            self._validate_footnote_references(master, footnote_definitions)
        footnote_ids = {
            note_id: index for index, note_id in enumerate(footnote_definitions, start=1)
        }
        bookmark_ids: dict[str, int] = {}
        if include_extras_only:
            document.add_heading("Дополнительные материалы", level=1)

        def add_table(table: MasterTable) -> None:
            nonlocal expected_tables
            if audio:
                document.add_heading(f"Смысл таблицы «{table.title}»", level=2)
                document.add_paragraph(table.audio_equivalent)
                return
            document.add_heading(table.title, level=2)
            word_table = document.add_table(rows=1, cols=len(table.headers))
            word_table.style = "Table Grid"
            for index, value in enumerate(table.headers):
                word_table.rows[0].cells[index].text = value
            for row in table.rows:
                cells = word_table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = value
            source_note = self._reader_source_note(table.source_note)
            if source_note:
                document.add_paragraph(f"Источник: {source_note}")
            expected_tables += 1

        def add_visual(visual: MasterVisual) -> None:
            nonlocal expected_visuals
            if audio:
                document.add_heading(f"Смысл материала «{visual.title}»", level=2)
                document.add_paragraph(visual.audio_equivalent)
                return
            image_path = self._render_visual(visual, visual_dir)
            document.add_picture(str(image_path), width=Inches(6.2))
            picture_paragraph = document.paragraphs[-1]
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph(f"{visual.caption}\nАльтернативный текст: {visual.alt_text}")
            source_note = self._reader_source_note(visual.source_note)
            if source_note:
                document.add_paragraph(f"Источник: {source_note}")
            expected_visuals += 1

        for chapter in master.chapters:
            if include_extras_only and not (chapter.tables or chapter.visuals):
                continue
            document.add_heading(chapter.title, level=1)
            if include_extras_only:
                for table in chapter.tables:
                    add_table(table)
                for visual in chapter.visuals:
                    add_visual(visual)
                continue
            for block_kind, block in self._ordered_chapter_blocks(chapter):
                if block_kind == "PARAGRAPH":
                    self._docx_add_semantic(document, str(block), footnote_ids, bookmark_ids)
                elif block_kind == "TABLE":
                    add_table(block)
                else:
                    add_visual(block)
        self._attach_docx_footnotes(document, footnote_definitions, footnote_ids)
        self._add_bibliography(document, master.bibliography)
        document.save(str(output))

        reopened = Document(str(output))
        embedded_visual_count = len(reopened.inline_shapes)
        footnote_part_present = not footnote_definitions
        document_xml = ""
        with ZipFile(output) as archive:
            footnote_part_present = (
                footnote_part_present or "word/footnotes.xml" in archive.namelist()
            )
            document_xml = archive.read("word/document.xml").decode("utf-8")
        qa = {
            "passed": (
                len(reopened.tables) == expected_tables
                and embedded_visual_count == expected_visuals
                and footnote_part_present
            ),
            "native_table_count": len(reopened.tables),
            "expected_table_count": expected_tables,
            "embedded_visual_count": embedded_visual_count,
            "expected_visual_count": expected_visuals,
            "footnote_definition_count": len(footnote_definitions),
            "native_footnote_part": footnote_part_present,
            "bookmark_count": document_xml.count("w:bookmarkStart"),
            "internal_hyperlink_count": document_xml.count("w:hyperlink"),
            "paragraph_count": len(reopened.paragraphs),
            "audio_adaptation": audio,
        }
        if not qa["passed"]:
            raise AutoBookExportError("DOCX structural/content QA failed")
        return qa

    def _audio_docx(self, script: AudioScriptView, output: Path, *, litres: bool) -> dict[str, Any]:
        if not script.ready_for_export:
            raise AudioScriptGateError(
                "audio DOCX requires a current human-approved AudioScript with no blocking checks"
            )
        document = Document()
        self._apply_docx_styles(document)
        document.add_heading(script.content.title, level=0)
        document.add_paragraph(script.content.author)
        for section in script.content.sections:
            document.add_heading(section.title, level=1)
            for paragraph in section.recording_paragraphs():
                document.add_paragraph(paragraph)
        document.save(str(output))
        reopened = Document(str(output))
        paragraph_count = len(reopened.paragraphs)
        expected = (
            2
            + len(script.content.sections)
            + sum(len(section.recording_paragraphs()) for section in script.content.sections)
        )
        if paragraph_count < expected:
            raise AutoBookExportError("AudioScript DOCX structural QA failed")
        return {
            "passed": True,
            "audio_script_id": script.audio_script_id,
            "audio_script_version": script.version,
            "audio_script_hash": script.content_hash,
            "source_hash": script.source_hash,
            "authority_status": script.status,
            "paragraph_count": paragraph_count,
            "same_snapshot_as_recording_txt": True,
            "platform_acceptance_claimed": False,
            "requirements_profile_checked_at": (
                "owner-litres-audio-profile.v1" if litres else "editorial-reading-profile.v1"
            ),
        }

    def _pdf(self, master: StructuredBookMaster, output: Path, visual_dir: Path) -> dict[str, Any]:
        footnote_definitions = self._footnote_definitions(master)
        self._validate_footnote_references(master, footnote_definitions)
        font = self._register_pdf_font()
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "BookOSBody", parent=styles["BodyText"], fontName=font, fontSize=11, leading=16
        )
        heading = ParagraphStyle(
            "BookOSHeading", parent=styles["Heading1"], fontName=font, fontSize=18, leading=23
        )
        title = ParagraphStyle(
            "BookOSTitle", parent=heading, fontSize=24, alignment=TA_CENTER, spaceAfter=18
        )
        story: list[Any] = [
            Paragraph(escape(master.title), title),
            Paragraph(escape(master.author), body),
            PageBreak(),
        ]
        for chapter in master.chapters:
            story.extend([Paragraph(escape(chapter.title), heading), Spacer(1, 4 * mm)])
            for block_kind, block in self._ordered_chapter_blocks(chapter):
                if block_kind == "PARAGRAPH":
                    story.extend(self._pdf_semantic_flowables(str(block), body, heading))
                elif block_kind == "TABLE":
                    table = block
                    story.append(Paragraph(escape(table.title), heading))
                    data = [[Paragraph(escape(cell), body) for cell in table.headers]]
                    data.extend(
                        [[Paragraph(escape(cell), body) for cell in row] for row in table.rows]
                    )
                    rendered = Table(data, repeatRows=1, hAlign="LEFT")
                    rendered.setStyle(
                        TableStyle(
                            [
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf0fb")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ]
                        )
                    )
                    story.extend([rendered, Spacer(1, 4 * mm)])
                    source_note = self._reader_source_note(table.source_note)
                    if source_note:
                        story.extend(
                            [Paragraph(escape(f"Источник: {source_note}"), body), Spacer(1, 2 * mm)]
                        )
                else:
                    visual = block
                    image_path = self._render_visual(visual, visual_dir)
                    story.extend(
                        [
                            PdfImage(str(image_path), width=160 * mm, height=90 * mm),
                            Paragraph(escape(visual.caption), body),
                        ]
                    )
                    source_note = self._reader_source_note(visual.source_note)
                    if source_note:
                        story.extend(
                            [Paragraph(escape(f"Источник: {source_note}"), body), Spacer(1, 2 * mm)]
                        )
        if master.bibliography:
            story.extend([PageBreak(), Paragraph("Библиография", heading)])
            story.extend(
                Paragraph(escape(f"{index}. {entry}"), body)
                for index, entry in enumerate(master.bibliography, start=1)
            )
        SimpleDocTemplate(
            str(output),
            pagesize=A4,
            rightMargin=20 * mm,
            leftMargin=20 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
            title=master.title,
            author=master.author,
        ).build(story)
        payload = output.read_bytes()
        reader = PdfReader(str(output))
        extracted_text = " ".join(
            " ".join((page.extract_text() or "").split()) for page in reader.pages
        )
        for note_id in footnote_definitions:
            extracted_text = extracted_text.replace(f"[{note_id}]", note_id)
        required_fragments = [master.title, *(chapter.title for chapter in master.chapters)]
        for chapter in master.chapters:
            for paragraph in chapter.paragraphs:
                required_fragments.extend(self._semantic_fragments(paragraph))
            for table in chapter.tables:
                required_fragments.extend([table.title, *table.headers, *table.rows[-1]])
            for visual in chapter.visuals:
                required_fragments.append(visual.caption)
        normalized_required = [
            " ".join(value.split()) for value in required_fragments if value.strip()
        ]
        content_preserved = all(value in extracted_text for value in normalized_required)
        passed = (
            payload.startswith(b"%PDF-")
            and b"%%EOF" in payload[-2048:]
            and bool(reader.pages)
            and content_preserved
        )
        if not passed:
            raise AutoBookExportError("PDF structural/content QA failed")
        return {
            "passed": True,
            "pdf_magic": True,
            "page_count": len(reader.pages),
            "reader_text_preserved": content_preserved,
            "visual_review_required": True,
        }

    def _epub(self, master: StructuredBookMaster, output: Path, visual_dir: Path) -> dict[str, Any]:
        self._validate_footnote_references(master, self._footnote_definitions(master))
        book = epub.EpubBook()
        book.set_identifier(master.manifest_hash)
        book.set_title(master.title)
        book.set_language(master.language)
        book.add_author(master.author)
        stylesheet = epub.EpubItem(
            uid="book-os-style",
            file_name="style/book.css",
            media_type="text/css",
            content=(
                b"body{line-height:1.45;} figure{text-align:center;} img{max-width:100%;height:auto;}"
                b"table{border-collapse:collapse;} th,td{border:1px solid #999;padding:.35em;}"
                b"aside.note,aside.footnote{margin:1em 0;padding:.6em;border-left:3px solid #777;}"
            ),
        )
        book.add_item(stylesheet)
        chapters: list[Any] = []
        for index, chapter in enumerate(master.chapters, start=1):
            item = epub.EpubHtml(
                title=chapter.title, file_name=f"chapter-{index}.xhtml", lang=master.language
            )
            item.add_link(href="style/book.css", rel="stylesheet", type="text/css")
            parts = [f"<h1>{escape(chapter.title)}</h1>"]
            for block_kind, block in self._ordered_chapter_blocks(chapter):
                if block_kind == "PARAGRAPH":
                    parts.append(self._epub_semantic_html(str(block)))
                elif block_kind == "TABLE":
                    table = block
                    parts.append(f"<h2>{escape(table.title)}</h2><table><thead><tr>")
                    parts.extend(f"<th>{escape(value)}</th>" for value in table.headers)
                    parts.append("</tr></thead><tbody>")
                    for row in table.rows:
                        parts.append(
                            "<tr>" + "".join(f"<td>{escape(value)}</td>" for value in row) + "</tr>"
                        )
                    parts.append("</tbody></table>")
                    source_note = self._reader_source_note(table.source_note)
                    if source_note:
                        parts.append(f'<p class="source-note">Источник: {escape(source_note)}</p>')
                else:
                    visual = block
                    image_path = self._render_visual(visual, visual_dir)
                    image_name = f"images/{visual.object_id}.png"
                    image_item = epub.EpubImage(
                        uid=f"visual-{visual.object_id}",
                        file_name=image_name,
                        media_type="image/png",
                        content=image_path.read_bytes(),
                    )
                    book.add_item(image_item)
                    parts.append(
                        f'<figure><img src="{escape(image_name)}" alt="{escape(visual.alt_text)}" />'
                        f"<figcaption>{escape(visual.caption)}</figcaption></figure>"
                    )
                    source_note = self._reader_source_note(visual.source_note)
                    if source_note:
                        parts.append(f'<p class="source-note">Источник: {escape(source_note)}</p>')
            item.content = "".join(parts)
            book.add_item(item)
            chapters.append(item)
        navigation: list[Any] = [*chapters]
        if master.bibliography:
            bibliography = epub.EpubHtml(
                title="Библиография",
                file_name="bibliography.xhtml",
                lang=master.language,
            )
            bibliography.add_link(href="style/book.css", rel="stylesheet", type="text/css")
            bibliography.content = (
                "<h1>Библиография</h1><ol>"
                + "".join(f"<li>{escape(entry)}</li>" for entry in master.bibliography)
                + "</ol>"
            )
            book.add_item(bibliography)
            navigation.append(bibliography)
        book.toc = tuple(navigation)
        book.spine = ["nav", *navigation]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(str(output), book, {})

        with ZipFile(output) as archive:
            if "mimetype" not in archive.namelist():
                raise AutoBookExportError("EPUB package has no mimetype")
            xhtml_names = [name for name in archive.namelist() if name.endswith(".xhtml")]
            for name in xhtml_names:
                try:
                    ET.fromstring(archive.read(name))
                except ET.ParseError as exc:
                    raise AutoBookExportError(f"EPUB XHTML is malformed: {name}") from exc
        reopened = epub.read_epub(str(output))
        document_count = len(list(reopened.get_items_of_type(ebooklib.ITEM_DOCUMENT)))
        expected_documents = len(master.chapters) + (1 if master.bibliography else 0)
        expected_visuals = sum(len(chapter.visuals) for chapter in master.chapters)
        image_count = len(list(reopened.get_items_of_type(ebooklib.ITEM_IMAGE)))
        document_items = list(reopened.get_items_of_type(ebooklib.ITEM_DOCUMENT))
        document_html = "\n".join(
            item.get_content().decode("utf-8", errors="replace") for item in document_items
        )
        expected_alt_texts = [
            visual.alt_text for chapter in master.chapters for visual in chapter.visuals
        ]
        expected_image_names = [
            f"images/{visual.object_id}.png"
            for chapter in master.chapters
            for visual in chapter.visuals
        ]
        references_ok = all(
            value in document_html for value in expected_alt_texts + expected_image_names
        )
        semantic_fragments = [
            fragment
            for chapter in master.chapters
            for paragraph in chapter.paragraphs
            for fragment in self._semantic_fragments(paragraph)
        ]
        semantic_documents: list[str] = []
        for item in document_items:
            try:
                root = ET.fromstring(item.get_content())
            except ET.ParseError as exc:
                raise AutoBookExportError("EPUB semantic QA could not parse a document") from exc
            semantic_documents.append("".join(root.itertext()))
        semantic_plain_text = " ".join(" ".join(semantic_documents).split())
        semantic_content_ok = all(
            " ".join(value.split()) in semantic_plain_text for value in semantic_fragments
        )
        if (
            document_count < expected_documents
            or image_count < expected_visuals
            or not references_ok
            or not semantic_content_ok
        ):
            raise AutoBookExportError("EPUB structural/content QA failed")
        return {
            "passed": True,
            "document_count": document_count,
            "expected_visual_count": expected_visuals,
            "image_count": image_count,
            "visual_references_and_alt_preserved": references_ok,
            "semantic_content_preserved": semantic_content_ok,
            "xhtml_parse_passed": True,
            "bibliography_document": bool(master.bibliography),
            "viewer_review_required": True,
        }

    @staticmethod
    def _master_as_pronunciation_content(master: StructuredBookMaster) -> AudioScriptContent:
        return AudioScriptContent(
            title=master.title,
            author=master.author,
            language=master.language,
            sections=[
                AudioScriptSection(
                    source_chapter_id=chapter.chapter_id,
                    title=chapter.title,
                    paragraphs=[
                        *chapter.paragraphs,
                        *(table.audio_equivalent for table in chapter.tables),
                        *(visual.audio_equivalent for visual in chapter.visuals),
                    ],
                )
                for chapter in master.chapters
            ],
        )

    @staticmethod
    def _publisher_pack(
        master: StructuredBookMaster,
        *,
        audit_bibliography: list[str],
        public_bibliography_included: bool,
    ) -> bytes:
        payload = {
            "title": master.title,
            "author": master.author,
            "language": master.language,
            "annotation": master.publisher_annotation,
            "chapter_count": len(master.chapters),
            "master_hash": master.manifest_hash,
            "public_bibliography_included": public_bibliography_included,
            "public_bibliography": master.bibliography,
            "bibliographic_audit": {
                "preserved_when_publicly_omitted": True,
                "verified_used_sources": audit_bibliography,
            },
        }
        return (json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()

    def export_selected(
        self,
        book_id: str,
        run_id: str,
        master: StructuredBookMaster,
        selection: AutoBookOutputSelection,
        *,
        audio_script: AudioScriptView | None = None,
        audit_bibliography: list[str] | None = None,
        public_bibliography_included: bool | None = None,
    ) -> ExportBundle:
        master_hash = master.manifest_hash
        output_dir = (
            self.runtime.projects.projects_dir / book_id / "exports" / run_id / master_hash[:12]
        )
        visual_dir = output_dir / "visuals"
        visual_dir.mkdir(parents=True, exist_ok=True)
        artifacts: list[AutoBookArtifactView] = []
        audio_service = AudioScriptService(self.data_dir)
        selected = selection.selected()
        if selection.audio_version_requested and audio_script is None:
            raise AudioScriptGateError(
                "audio outputs cannot be created directly from Literary Master; "
                "a human-approved AudioScript is required"
            )
        if audio_script is not None and not audio_script.ready_for_export:
            raise AudioScriptGateError("the selected AudioScript is not approved and current")
        voice_payload: bytes | None = None
        voice_relative_path: str | None = None
        internal_bibliography = (
            list(master.bibliography) if audit_bibliography is None else list(audit_bibliography)
        )
        public_included = (
            bool(master.bibliography)
            if public_bibliography_included is None
            else public_bibliography_included
        )
        for kind in selected:
            audio_bound = kind in {
                "AUDIO_READING_DOCX",
                "AUDIO_LITRES_DOCX",
                "VOICE_TEXT_TXT",
                "AUDIO_PRODUCTION_HANDOFF",
            } or (kind == "PRONUNCIATION_DICTIONARY" and audio_script is not None)
            kind_dir = (
                output_dir / f"audio-{audio_script.content_hash[:12]}"
                if audio_bound and audio_script is not None
                else output_dir
            )
            kind_dir.mkdir(parents=True, exist_ok=True)
            output = kind_dir / self._FILE_NAMES[kind]
            if kind == "FULL_MANUSCRIPT_DOCX":
                qa = self._docx(master, output, audio=False, visual_dir=visual_dir)
            elif kind == "LITRES_EBOOK_DOCX":
                qa = self._docx(master, output, audio=False, visual_dir=visual_dir)
                qa["platform_acceptance_claimed"] = False
                qa["requirements_profile"] = self.LITRES_PROFILE_VERSION
                qa["requirements_profile_checked_at"] = self.LITRES_PROFILE_CHECKED_AT
                qa["embedded_images_rgb"] = True
                qa["embedded_image_max_long_side_px"] = 1024
                qa["embedded_image_max_area_px"] = 2_000_000
                qa["image_layout"] = "inline-with-text"
                qa["file_size_bytes"] = output.stat().st_size
                qa["within_70_mb_upload_limit"] = output.stat().st_size <= 70 * 1024 * 1024
            elif kind in {"AUDIO_READING_DOCX", "AUDIO_LITRES_DOCX"}:
                assert audio_script is not None
                qa = self._audio_docx(
                    audio_script,
                    output,
                    litres=kind == "AUDIO_LITRES_DOCX",
                )
            elif kind == "READING_PDF":
                qa = self._pdf(master, output, visual_dir)
            elif kind == "EPUB":
                qa = self._epub(master, output, visual_dir)
            elif kind == "VOICE_TEXT_TXT":
                assert audio_script is not None
                voice_payload = audio_script.content.clean_recording_text().encode("utf-8")
                output.write_bytes(voice_payload)
                voice_relative_path = str(
                    output.relative_to(self.runtime.projects.projects_dir / book_id)
                )
                qa = {
                    "passed": True,
                    "encoding": "UTF-8",
                    "clean_recording_text": True,
                    "audio_script_id": audio_script.audio_script_id,
                    "audio_script_hash": audio_script.content_hash,
                    "source_hash": audio_script.source_hash,
                    "contains_ssml": False,
                    "contains_technical_metadata": False,
                }
            elif kind == "PRONUNCIATION_DICTIONARY":
                content = (
                    audio_script.content
                    if audio_script is not None
                    else self._master_as_pronunciation_content(master)
                )
                entries = (
                    audio_script.pronunciation_entries
                    if audio_script is not None
                    else audio_service.discover_pronunciation(content)
                )
                output.write_text(audio_service.pronunciation_text(entries), encoding="utf-8")
                qa = {
                    "passed": True,
                    "technical_dictionary_outside_manuscript": True,
                    "entry_count": len(entries),
                    "unverified_entries": sum(item.status == "NEEDS_REVIEW" for item in entries),
                    "empty_recommendations_are_not_verified": all(
                        item.recommendation or item.status != "VERIFIED" for item in entries
                    ),
                }
            elif kind == "AUDIO_PRODUCTION_HANDOFF":
                assert audio_script is not None
                if voice_payload is None or voice_relative_path is None:
                    raise AudioScriptGateError("handoff requires the mandatory UTF-8 recording TXT")
                manifest = audio_service.record_handoff(
                    book_id,
                    audio_script,
                    text_relative_path=voice_relative_path,
                    text_payload=voice_payload,
                )
                handoff_payload = (
                    json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
                ).encode("utf-8")
                handoff_hash = hashlib.sha256(handoff_payload).hexdigest()
                existing_handoff = next(
                    (
                        item
                        for item in self.runtime.list_artifacts(book_id, run_id)
                        if item.output_kind == "AUDIO_PRODUCTION_HANDOFF"
                        and item.master_hash == audio_script.content_hash
                        and item.status == "READY"
                    ),
                    None,
                )
                if existing_handoff is not None:
                    expected_relative_path = str(
                        output.relative_to(self.runtime.projects.projects_dir / book_id)
                    )
                    if (
                        existing_handoff.relative_path != expected_relative_path
                        or existing_handoff.content_hash != handoff_hash
                    ):
                        raise AudioScriptGateError(
                            "the immutable handoff payload conflicts with its registered artifact"
                        )
                    if output.exists() and output.read_bytes() != handoff_payload:
                        raise AudioScriptGateError(
                            "the immutable handoff file differs from its registered artifact"
                        )
                if not output.exists():
                    output.write_bytes(handoff_payload)
                elif existing_handoff is None:
                    output.write_bytes(handoff_payload)
                qa = {
                    "passed": True,
                    "schema": "book-os-audiobook-handoff.v2",
                    "audio_script_hash": audio_script.content_hash,
                    "recording_text_hash": manifest["recording_text"]["content_hash"],
                }
            elif kind == "READER_EXTRAS":
                if not (
                    any(chapter.tables or chapter.visuals for chapter in master.chapters)
                    or master.bibliography
                ):
                    raise AutoBookExportError(
                        "selected reader extras contain no reader-facing tables, visuals, or bibliography"
                    )
                qa = self._docx(
                    master, output, audio=False, include_extras_only=True, visual_dir=visual_dir
                )
            else:
                output.write_bytes(
                    self._publisher_pack(
                        master,
                        audit_bibliography=internal_bibliography,
                        public_bibliography_included=public_included,
                    )
                )
                qa = {
                    "passed": True,
                    "schema": "publisher-pack.v1",
                    "public_bibliography_included": public_included,
                    "bibliographic_audit_source_count": len(internal_bibliography),
                }
            payload = output.read_bytes()
            relative_path = str(output.relative_to(self.runtime.projects.projects_dir / book_id))
            artifacts.append(
                self.runtime.register_artifact(
                    book_id,
                    run_id,
                    output_kind=kind,
                    master_hash=(
                        audio_script.content_hash
                        if audio_bound and audio_script is not None
                        else master_hash
                    ),
                    profile_version=self.PROFILE_VERSIONS[kind],
                    exporter_version=self.EXPORTER_VERSION,
                    relative_path=relative_path,
                    payload=payload,
                    qa=qa,
                )
            )
        return ExportBundle(
            master_hash=master_hash,
            artifacts=artifacts,
            output_directory=str(output_dir),
        )
