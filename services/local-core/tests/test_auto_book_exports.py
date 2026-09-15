from pathlib import Path
import hashlib
import json
import sqlite3
import pytest
from zipfile import ZipFile

from docx import Document
from PIL import Image
from pypdf import PdfReader
import ebooklib  # type: ignore[import-untyped]
from ebooklib import epub

from book_os_core.auto_book_exports import (
    AutoBookExportError,
    AutoBookExporter,
    MasterChapter,
    MasterTable,
    MasterVisual,
    StructuredBookMaster,
)
from book_os_core.audio_script import AudioScriptService
from book_os_core.auto_book_runtime import (
    AutoBookIntent,
    AutoBookOutputSelection,
    DurableAutoBookRuntime,
)
from book_os_core.projects import NewBookRequest, ProjectService


def setup_run(tmp_path: Path, outputs: AutoBookOutputSelection) -> tuple[str, str]:
    book_id = (
        ProjectService(tmp_path)
        .create_project(NewBookRequest(working_title="Fixture export", primary_subtype="Strategy"))
        .book_id
    )
    run = DurableAutoBookRuntime(tmp_path).create_run(
        book_id,
        AutoBookIntent(
            idea="Проверить один master и независимые производные результаты.",
            author_name="Елена Дым",
            outputs=outputs,
            max_total_cost_usd=5,
            max_cost_usd_per_request=1,
            max_requests=10,
        ),
    )
    return book_id, run.run_id


def fixture_master(*, suffix: str = "") -> StructuredBookMaster:
    tables = [
        MasterTable(
            object_id=f"table-{index}",
            title=f"Таблица {index}",
            headers=["Шаг", "Результат"],
            rows=[["Проверить", f"Получить результат {index}"]],
            audio_equivalent=(
                f"Таблица {index} показывает: сначала нужно проверить условие, "
                f"после чего получить результат {index}."
            ),
        )
        for index in range(1, 17)
    ]
    return StructuredBookMaster(
        title=f"Проверяемая книга{suffix}",
        author="Елена Дым",
        chapters=[
            MasterChapter(
                chapter_id="chapter-1",
                title="Глава 1. Решение",
                paragraphs=[
                    "Сильный практический текст объясняет механизм, границы и следующий шаг.",
                    "Каждый смысловой объект сохраняется во всех подходящих форматах.",
                ],
                tables=tables,
                visuals=[
                    MasterVisual(
                        object_id="chart-1",
                        kind="CHART",
                        title="Динамика результата",
                        caption="Рисунок 1. Условные данные для проверки экспорта.",
                        alt_text="Три столбца растут от одного до трёх.",
                        audio_equivalent=(
                            "График на условных данных показывает последовательный рост "
                            "показателя от 1 до 3; для записи числа нужно произнести словами."
                        ),
                        data=[("А", 1), ("Б", 2), ("В", 3)],
                        source_note="Условные fixture-данные, не фактическое утверждение",
                    )
                ],
            )
        ],
        bibliography=["Тестовый источник. URL: https://example.test/source"],
        publisher_annotation="Книга о проверяемом производственном цикле.",
    )


def test_selected_outputs_keep_sixteen_native_tables_and_audio_meaning(tmp_path: Path) -> None:
    selection = AutoBookOutputSelection(
        full_manuscript_docx=True,
        litres_ebook_docx=True,
        reading_pdf=True,
        epub=True,
        audio_reading_docx=True,
        audio_litres_docx=True,
        voice_text_txt=True,
        pronunciation_dictionary=True,
        reader_extras=True,
        publisher_pack=True,
    )
    book_id, run_id = setup_run(tmp_path, selection)
    master = fixture_master()
    audio_service = AudioScriptService(tmp_path)
    content, transformations = audio_service.content_from_master(
        master,
        adaptation_mode="SOURCE_FAITHFUL",
        adapted_chapters={
            "chapter-1": (
                "Сначала разберём механизм решения так, чтобы его можно было понять с первого "
                "прослушивания.\n\nЗатем свяжем проверку с практическим следующим шагом."
            )
        },
    )
    proposed = audio_service.create_proposal(
        book_id,
        source_kind="LITERARY_MASTER",
        source_identity="master-fixture-v1",
        source_hash=master.manifest_hash,
        adaptation_mode="SOURCE_FAITHFUL",
        content=content,
        transformations=transformations,
        provenance={
            "operation": "fixture audio edit",
            "provider_calls": 0,
            "source_attribution": master.bibliography,
        },
    )
    attention = sorted(
        {
            finding.code
            for check in proposed.quality_checks
            for finding in check.findings
            if finding.severity == "ATTENTION"
        }
    )
    approved = audio_service.approve(
        book_id,
        proposed.audio_script_id,
        human_actor="Owner fixture",
        accepted_attention_codes=attention,
    )
    bundle = AutoBookExporter(tmp_path, DurableAutoBookRuntime(tmp_path)).export_selected(
        book_id,
        run_id,
        master,
        selection,
        audio_script=approved,
    )

    assert len(bundle.artifacts) == 11
    by_kind = {artifact.output_kind: artifact for artifact in bundle.artifacts}
    project_dir = tmp_path / "projects" / book_id

    manuscript = Document(project_dir / by_kind["FULL_MANUSCRIPT_DOCX"].relative_path)
    assert len(manuscript.tables) == 16
    assert "Библиография" in [paragraph.text for paragraph in manuscript.paragraphs]

    audio = Document(project_dir / by_kind["AUDIO_READING_DOCX"].relative_path)
    assert len(audio.tables) == 0
    audio_text = "\n".join(paragraph.text for paragraph in audio.paragraphs)
    assert "Таблица 16 показывает" in audio_text
    assert "График на условных данных" in audio_text

    voice = (project_dir / by_kind["VOICE_TEXT_TXT"].relative_path).read_text(encoding="utf-8")
    assert "Таблица 1 показывает" in voice
    assert "График на условных данных" in voice
    assert "|" not in voice
    assert "Библиография доступна" not in voice
    assert by_kind["VOICE_TEXT_TXT"].qa["audio_script_hash"] == approved.content_hash
    assert by_kind["AUDIO_READING_DOCX"].qa["audio_script_hash"] == approved.content_hash

    handoff = (project_dir / by_kind["AUDIO_PRODUCTION_HANDOFF"].relative_path).read_text(
        encoding="utf-8"
    )
    assert approved.audio_script_id in handoff
    assert approved.source_hash in handoff
    assert master.bibliography[0] in handoff
    first_handoff_artifact = by_kind["AUDIO_PRODUCTION_HANDOFF"]
    first_handoff_bytes = (project_dir / first_handoff_artifact.relative_path).read_bytes()

    repeated = AutoBookExporter(tmp_path, DurableAutoBookRuntime(tmp_path)).export_selected(
        book_id,
        run_id,
        master,
        selection,
        audio_script=approved,
    )
    repeated_handoff_artifact = next(
        item for item in repeated.artifacts if item.output_kind == "AUDIO_PRODUCTION_HANDOFF"
    )
    repeated_handoff_bytes = (project_dir / repeated_handoff_artifact.relative_path).read_bytes()
    first_manifest = json.loads(first_handoff_bytes)
    repeated_manifest = json.loads(repeated_handoff_bytes)
    assert repeated_handoff_artifact.artifact_id == first_handoff_artifact.artifact_id
    assert repeated_handoff_artifact.content_hash == first_handoff_artifact.content_hash
    assert repeated_handoff_bytes == first_handoff_bytes
    assert repeated_manifest["handoff_id"] == first_manifest["handoff_id"]
    assert (
        hashlib.sha256(repeated_handoff_bytes).hexdigest() == repeated_handoff_artifact.content_hash
    )
    with sqlite3.connect(project_dir / "project.sqlite") as connection:
        persisted = connection.execute(
            "SELECT handoff_id,manifest_json FROM audio_production_handoffs "
            "WHERE audio_script_id=? AND script_hash=?",
            (approved.audio_script_id, approved.content_hash),
        ).fetchall()
    assert len(persisted) == 1
    assert persisted[0][0] == first_manifest["handoff_id"]
    assert json.loads(persisted[0][1]) == first_manifest

    pdf_payload = (project_dir / by_kind["READING_PDF"].relative_path).read_bytes()
    assert pdf_payload.startswith(b"%PDF-")
    with ZipFile(project_dir / by_kind["EPUB"].relative_path) as archive:
        assert "mimetype" in archive.namelist()
        bibliography_files = [name for name in archive.namelist() if "bibliography" in name]
        assert len(bibliography_files) == 1
        assert "Тестовый источник" in archive.read(bibliography_files[0]).decode("utf-8")

    extras = Document(project_dir / by_kind["READER_EXTRAS"].relative_path)
    assert "Библиография" in [paragraph.text for paragraph in extras.paragraphs]
    publisher = json.loads(
        (project_dir / by_kind["PUBLISHER_PACK"].relative_path).read_text(encoding="utf-8")
    )
    assert publisher["public_bibliography_included"] is True
    assert publisher["public_bibliography"] == master.bibliography
    assert publisher["bibliographic_audit"]["verified_used_sources"] == master.bibliography

    visual = Path(bundle.output_directory) / "visuals" / "chart-1.png"
    assert visual.is_file()
    assert visual.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
    assert all(artifact.qa["passed"] is True for artifact in bundle.artifacts)


def test_public_opt_out_hides_reader_bibliography_but_keeps_publisher_audit(
    tmp_path: Path,
) -> None:
    selection = AutoBookOutputSelection(
        full_manuscript_docx=True,
        epub=True,
        reader_extras=True,
        publisher_pack=True,
    )
    book_id, run_id = setup_run(tmp_path, selection)
    public_master = fixture_master()
    internal_sources = list(public_master.bibliography)
    omitted_master = public_master.model_copy(update={"bibliography": []})
    bundle = AutoBookExporter(tmp_path, DurableAutoBookRuntime(tmp_path)).export_selected(
        book_id,
        run_id,
        omitted_master,
        selection,
        audit_bibliography=internal_sources,
        public_bibliography_included=False,
    )
    by_kind = {artifact.output_kind: artifact for artifact in bundle.artifacts}
    project_dir = tmp_path / "projects" / book_id
    manuscript = Document(project_dir / by_kind["FULL_MANUSCRIPT_DOCX"].relative_path)
    extras = Document(project_dir / by_kind["READER_EXTRAS"].relative_path)
    assert "Библиография" not in [paragraph.text for paragraph in manuscript.paragraphs]
    assert "Библиография" not in [paragraph.text for paragraph in extras.paragraphs]
    with ZipFile(project_dir / by_kind["EPUB"].relative_path) as archive:
        assert not [name for name in archive.namelist() if "bibliography" in name]
    publisher = json.loads(
        (project_dir / by_kind["PUBLISHER_PACK"].relative_path).read_text(encoding="utf-8")
    )
    assert publisher["public_bibliography"] == []
    assert publisher["public_bibliography_included"] is False
    assert publisher["bibliographic_audit"]["verified_used_sources"] == internal_sources


def test_late_output_selection_exports_only_missing_derivative_and_marks_old_stale(
    tmp_path: Path,
) -> None:
    initial = AutoBookOutputSelection(full_manuscript_docx=True)
    book_id, run_id = setup_run(tmp_path, initial)
    runtime = DurableAutoBookRuntime(tmp_path)
    exporter = AutoBookExporter(tmp_path, runtime)
    first = exporter.export_selected(book_id, run_id, fixture_master(), initial)
    assert [item.output_kind for item in first.artifacts] == ["FULL_MANUSCRIPT_DOCX"]

    later = AutoBookOutputSelection(full_manuscript_docx=False, epub=True)
    second = exporter.export_selected(book_id, run_id, fixture_master(), later)
    assert [item.output_kind for item in second.artifacts] == ["EPUB"]
    assert len(runtime.list_artifacts(book_id, run_id)) == 2

    changed = AutoBookOutputSelection(full_manuscript_docx=True)
    exporter.export_selected(book_id, run_id, fixture_master(suffix=" — новая ревизия"), changed)
    manuscripts = [
        item
        for item in runtime.list_artifacts(book_id, run_id)
        if item.output_kind == "FULL_MANUSCRIPT_DOCX"
    ]
    assert [item.status for item in manuscripts] == ["STALE", "READY"]


def test_new_audio_script_version_preserves_old_files_and_marks_old_artifacts_stale(
    tmp_path: Path,
) -> None:
    selection = AutoBookOutputSelection(
        full_manuscript_docx=False,
        audio_reading_docx=True,
    )
    book_id, run_id = setup_run(tmp_path, selection)
    master = fixture_master()
    service = AudioScriptService(tmp_path)
    exporter = AutoBookExporter(tmp_path, DurableAutoBookRuntime(tmp_path))
    paths: list[str] = []
    for suffix in ("Первая аудиоредакция.", "Вторая исправленная аудиоредакция."):
        content, mapping = service.content_from_master(
            master,
            adaptation_mode="SOURCE_FAITHFUL",
            adapted_chapters={"chapter-1": suffix},
        )
        proposed = service.create_proposal(
            book_id,
            source_kind="LITERARY_MASTER",
            source_identity="master-fixture-v1",
            source_hash=master.manifest_hash,
            adaptation_mode="SOURCE_FAITHFUL",
            content=content,
            transformations=mapping,
            provenance={"provider_calls": 0},
        )
        attention = sorted(
            {
                finding.code
                for check in proposed.quality_checks
                for finding in check.findings
                if finding.severity == "ATTENTION"
            }
        )
        approved = service.approve(
            book_id,
            proposed.audio_script_id,
            human_actor="Owner",
            accepted_attention_codes=attention,
        )
        bundle = exporter.export_selected(
            book_id,
            run_id,
            master,
            selection,
            audio_script=approved,
        )
        paths.append(
            next(
                item.relative_path
                for item in bundle.artifacts
                if item.output_kind == "AUDIO_READING_DOCX"
            )
        )

    assert paths[0] != paths[1]
    project = tmp_path / "projects" / book_id
    assert all((project / item).is_file() for item in paths)
    audio_artifacts = [
        item
        for item in DurableAutoBookRuntime(tmp_path).list_artifacts(book_id, run_id)
        if item.output_kind == "AUDIO_READING_DOCX"
    ]
    assert [item.status for item in audio_artifacts] == ["STALE", "READY"]


def test_ordered_blocks_preserve_between_paragraph_placement() -> None:
    chapter = MasterChapter(
        chapter_id="c1",
        title="Глава",
        paragraphs=["Первый абзац.", "Второй абзац.", "Третий абзац."],
        tables=[
            MasterTable(
                object_id="t1",
                title="Таблица",
                headers=["A", "B"],
                rows=[["1", "2"]],
                audio_equivalent="A один, B два.",
                placement_after_paragraph=1,
            )
        ],
        visuals=[
            MasterVisual(
                object_id="v1",
                kind="CHART",
                title="Диаграмма",
                caption="Подпись",
                alt_text="Столбец один.",
                audio_equivalent="Значение один.",
                data=[("Один", 1.0)],
                placement_after_paragraph=2,
            )
        ],
    )
    blocks = AutoBookExporter._ordered_chapter_blocks(chapter)
    assert [(kind, getattr(value, "object_id", value)) for kind, value in blocks] == [
        ("PARAGRAPH", "Первый абзац."),
        ("TABLE", "t1"),
        ("PARAGRAPH", "Второй абзац."),
        ("VISUAL", "v1"),
        ("PARAGRAPH", "Третий абзац."),
    ]


def test_epub_physically_contains_visual_and_alt_text(tmp_path: Path) -> None:
    exporter = AutoBookExporter(tmp_path, DurableAutoBookRuntime(tmp_path))
    master = StructuredBookMaster(
        title="Книга",
        author="Автор",
        chapters=[
            MasterChapter(
                chapter_id="c1",
                title="Глава",
                paragraphs=["До изображения.", "После изображения."],
                visuals=[
                    MasterVisual(
                        object_id="visual-one",
                        kind="CHART",
                        title="Диаграмма",
                        caption="Проверенная подпись",
                        alt_text="Альтернативное описание диаграммы",
                        audio_equivalent="Диаграмма показывает значение десять.",
                        data=[("Показатель", 10.0)],
                        placement_after_paragraph=1,
                    )
                ],
            )
        ],
    )
    output = tmp_path / "visual.epub"
    qa = exporter._epub(master, output, tmp_path / "visuals")
    assert qa["expected_visual_count"] == 1
    assert qa["image_count"] >= 1
    reopened = epub.read_epub(str(output))
    chapter = next(iter(reopened.get_items_of_type(ebooklib.ITEM_DOCUMENT)))
    content = chapter.get_content().decode("utf-8")
    assert "Альтернативное описание диаграммы" in content
    assert (
        content.index("До изображения")
        < content.index("visual-one.png")
        < content.index("После изображения")
    )


def test_semantic_text_long_table_native_footnote_and_crossref_survive_exports(
    tmp_path: Path,
) -> None:
    exporter = AutoBookExporter(tmp_path, DurableAutoBookRuntime(tmp_path))
    long_rows = [[str(index), f"Строка {index}"] for index in range(1, 121)]
    master = StructuredBookMaster(
        title="Семантическая книга",
        author="Автор",
        chapters=[
            MasterChapter(
                chapter_id="c-semantic",
                title="Глава семантики",
                paragraphs=[
                    "## Решение {#decision}",
                    "Обычный **жирный** и *курсивный* текст со сноской[^1].",
                    "- Первый пункт\n- Второй пункт",
                    "1. Первый шаг\n2. Второй шаг",
                    "> Важное примечание для читателя.",
                    "См. [раздел решения](#decision).",
                    "[^1]: Смысл сноски должен сохраниться.",
                ],
                tables=[
                    MasterTable(
                        object_id="long-table",
                        title="Длинная таблица",
                        headers=["N", "Значение"],
                        rows=long_rows,
                        audio_equivalent="Таблица перечисляет сто двадцать строк.",
                        source_note="Проверяемая таблица автора",
                        placement_after_paragraph=4,
                    )
                ],
                visuals=[
                    MasterVisual(
                        object_id="semantic-chart",
                        kind="CHART",
                        title="Диаграмма",
                        caption="Проверенная диаграмма",
                        alt_text="Диаграмма с одним значением десять.",
                        audio_equivalent="Диаграмма показывает значение десять.",
                        data=[("Показатель", 10.0)],
                        source_note="Проверяемые данные автора",
                        placement_after_paragraph=5,
                    )
                ],
            )
        ],
        bibliography=["Проверяемый источник. https://example.test/source"],
    )
    visual_dir = tmp_path / "visuals"
    visual_dir.mkdir()
    docx_path = tmp_path / "book.docx"
    pdf_path = tmp_path / "book.pdf"
    epub_path = tmp_path / "book.epub"
    docx_qa = exporter._docx(master, docx_path, audio=False, visual_dir=visual_dir)
    pdf_qa = exporter._pdf(master, pdf_path, visual_dir)
    epub_qa = exporter._epub(master, epub_path, visual_dir)
    assert docx_qa["embedded_visual_count"] == 1
    assert docx_qa["native_footnote_part"] is True
    assert docx_qa["bookmark_count"] >= 1
    assert docx_qa["internal_hyperlink_count"] >= 1
    assert pdf_qa["reader_text_preserved"] is True
    assert epub_qa["semantic_content_preserved"] is True
    assert epub_qa["visual_references_and_alt_preserved"] is True
    assert epub_qa["xhtml_parse_passed"] is True

    doc = Document(docx_path)
    doc_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Решение" in doc_text
    assert "Первый пункт" in doc_text and "Второй шаг" in doc_text
    assert "Важное примечание" in doc_text
    assert any(
        run.bold and run.text == "жирный" for paragraph in doc.paragraphs for run in paragraph.runs
    )
    assert any(
        run.italic and run.text == "курсивный"
        for paragraph in doc.paragraphs
        for run in paragraph.runs
    )
    assert doc.tables[0].rows[-1].cells[1].text == "Строка 120"
    with ZipFile(docx_path) as archive:
        assert "word/footnotes.xml" in archive.namelist()
        footnotes = archive.read("word/footnotes.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        assert "Смысл сноски должен сохраниться" in footnotes
        assert 'w:anchor="decision"' in document_xml
        assert 'w:name="decision"' in document_xml

    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages)
    assert "Первый пункт" in pdf_text
    assert "Строка 120" in pdf_text
    assert "Смысл сноски должен сохраниться" in pdf_text
    assert "Источник: Проверяемая таблица автора" in pdf_text

    reopened = epub.read_epub(str(epub_path))
    html = "\n".join(
        item.get_content().decode("utf-8", errors="replace")
        for item in reopened.get_items_of_type(ebooklib.ITEM_DOCUMENT)
    )
    assert '<h2 id="decision">Решение</h2>' in html
    assert "<ul>" in html and "<ol>" in html
    assert "<strong>жирный</strong>" in html and "<em>курсивный</em>" in html
    assert 'href="#decision"' in html
    assert 'id="fn-1"' in html
    assert 'href="#fn-1"' in html
    assert "Строка 120" in html
    assert "Источник: Проверяемые данные автора" in html

    visual_path = visual_dir / "semantic-chart.png"
    with Image.open(visual_path) as image:
        assert image.mode == "RGB"
        assert max(image.size) <= 1024
        assert image.width * image.height <= 2_000_000


def test_reader_extras_refuses_empty_selected_derivative(tmp_path: Path) -> None:
    selection = AutoBookOutputSelection(reader_extras=True)
    book_id, run_id = setup_run(tmp_path, selection)
    master = StructuredBookMaster(
        title="Книга без приложений",
        author="Автор",
        chapters=[MasterChapter(chapter_id="c1", title="Глава", paragraphs=["Текст."])],
    )
    with pytest.raises(AutoBookExportError, match="no reader-facing"):
        AutoBookExporter(tmp_path, DurableAutoBookRuntime(tmp_path)).export_selected(
            book_id, run_id, master, selection
        )
