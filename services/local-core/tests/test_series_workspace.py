import base64
from io import BytesIO
import json
from pathlib import Path

from docx import Document
import pytest
from sqlalchemy import text

from book_os_core.book_context import ProfileCreateRequest, ProfileRegistry
from book_os_core.db import create_database
from book_os_core.projects import BookArchitecturePayload, BookContractPayload, ProjectService
from book_os_core.series_workspace import (
    SeriesBookCreateRequest,
    SeriesCreateRequest,
    SeriesExportSelection,
    SeriesImportRequest,
    SeriesPresetRequest,
    SeriesWorkspaceGateError,
    SeriesWorkspaceService,
)


def approved_author(tmp_path: Path) -> str:
    registry = ProfileRegistry(tmp_path)
    profile = registry.create_profile(
        ProfileCreateRequest(kind="AUTHOR", content={"author_name": "Елена Дым"})
    )
    return registry.approve_profile(profile.profile_id).profile_id


def new_series(tmp_path: Path) -> tuple[SeriesWorkspaceService, str]:
    service = SeriesWorkspaceService(tmp_path)
    profile = service.create_series(
        SeriesCreateRequest(
            series_name="Сильная серия",
            author_profile_id=approved_author(tmp_path),
            audience="Авторы практического нон-фикшн",
            promise="Каждая книга решает отдельную большую задачу",
            territory="Профессиональная работа с книгой",
        )
    )
    ProfileRegistry(tmp_path).approve_profile(profile.profile_id)
    return service, profile.profile_id


def book(title: str, ordinal: int, idea: str) -> SeriesBookCreateRequest:
    return SeriesBookCreateRequest(
        title=title,
        ordinal=ordinal,
        unique_idea=idea,
        reader_problem=f"Проблема для {title}",
        reader_result=f"Самостоятельный результат {title}",
        unique_mechanism=f"Уникальный механизм {title}",
        excluded_topics=["Территория соседней книги"],
    )


def test_series_map_is_persistent_owner_approved_and_becomes_stale(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    first = service.add_book(series_id, book("Архитектура", 1, "Построить неповторимую структуру"))
    second = service.add_book(series_id, book("Редактура", 2, "Исправить сквозные дефекты текста"))
    service.approve_book_passport(series_id, first.book_id, first.passport_hash, "Утверждаю")
    service.approve_book_passport(series_id, second.book_id, second.passport_hash, "Утверждаю")

    result = service.analyze(series_id)
    assert result.status in {"PASS", "ATTENTION"}
    assert not result.approved
    approved = service.approve_map(series_id, result.map_hash, "Границы книг проверены")
    assert approved.approved
    service.require_current_map(series_id)

    service.add_book(series_id, book("Факты", 3, "Проверить доказательную основу"))
    stale = service.current_map(series_id)
    assert stale is not None and not stale.current
    with pytest.raises(SeriesWorkspaceGateError, match="Book Passports|missing or stale"):
        service.require_current_map(series_id)


def test_import_keeps_original_bytes_and_a_separate_structural_profile(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    entry = service.add_book(
        series_id,
        book("Внешняя книга", 1, "Разобрать уже написанную авторскую рукопись"),
    )
    document = Document()
    document.add_heading("Глава первая", level=1)
    document.add_paragraph("Точный исходный текст не переписывается при импорте.")
    document.add_table(rows=2, cols=2)
    stream = BytesIO()
    document.save(stream)
    payload = stream.getvalue()

    imported = service.import_source(
        series_id,
        entry.book_id,
        SeriesImportRequest(
            filename="не-считать-названием.docx",
            content_base64=base64.b64encode(payload).decode(),
            rights_status="AUTHOR_MANUSCRIPT",
        ),
    )

    assert imported.analysis_status == "PARSED"
    assert imported.analysis["filename_is_not_book_title"] is True
    assert imported.analysis["headings"] == ["Глава первая"]
    assert imported.analysis["tables"] == 1
    stored = tmp_path / "projects" / entry.book_id / "series-imports"
    assert next(stored.iterdir()).read_bytes() == payload


def test_semantic_overlap_blocks_owner_map_approval(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    repeated = "Одна книга объясняет центральный механизм продаж услуг через доверие клиента"
    service.add_book(series_id, book("Книга один", 1, repeated))
    service.add_book(series_id, book("Книга два", 2, repeated))

    result = service.analyze(series_id)
    assert result.status == "BLOCKING"
    with pytest.raises(SeriesWorkspaceGateError, match="blocking overlap"):
        service.approve_map(series_id, result.map_hash, "Одобряю")


def test_current_architecture_clone_is_a_persisted_pre_writing_blocker(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    first = service.add_book(
        series_id,
        book("Первая территория", 1, "Диагностировать отдельную проблему спроса"),
    )
    second = service.add_book(
        series_id,
        book("Вторая территория", 2, "Настроить отдельную систему удержания"),
    )
    projects = ProjectService(tmp_path)
    contract = BookContractPayload(
        reader="Владелец профессиональной практики",
        reader_problem="Неясно, как устроен отдельный механизм книги",
        central_promise="Получить проверяемую модель принятия решений",
        central_thesis="Система становится управляемой через явные критерии",
        unique_angle="Разобрать причинную конструкцию на наблюдаемых решениях",
        reader_trajectory="От симптомов к самостоятельной диагностике",
        explicit_exclusions=["Не каталог общих советов"],
        evidence_policy="Материальные утверждения требуют evidence",
        voice_genre_constraints="Точный практический нон-фикшн",
        readiness_criteria=["Читатель может применить критерий"],
    )
    cloned_architecture = BookArchitecturePayload.model_validate(
        {
            "parts": [
                {
                    "title": "Причинная система",
                    "purpose": "Собрать последовательность решений",
                    "chapters": [
                        {
                            "title": "Диагностика ограничения",
                            "purpose": "Найти наблюдаемую причину ограничения",
                            "new_contribution": "Карта причин и проверяемых последствий",
                        },
                        {
                            "title": "Перестройка решения",
                            "purpose": "Перевести диагноз в новое правило действия",
                            "new_contribution": "Контур внедрения и обратной связи",
                            "dependencies": ["Диагностика ограничения"],
                            "transition": "От причины к изменению системы",
                        },
                    ],
                }
            ],
            "intellectual_progression": "Диагноз переходит в устройство решения",
            "concept_allocation": "Главы владеют разными причинными функциями",
            "promise_thesis_coverage": "Вся последовательность выполняет обещание",
            "major_transitions": "Причина открывает способ изменения",
        }
    )
    for entry in (first, second):
        projects.save_book_contract(entry.book_id, contract)
        projects.approve_book_contract(entry.book_id)
        projects.save_architecture(entry.book_id, cloned_architecture)
        projects.approve_architecture(entry.book_id)

    result = service.analyze(series_id)
    assert result.status == "BLOCKING"
    architecture = [
        item
        for item in result.findings
        if item["dimension"] == "ARCHITECTURE" and "identical_chapter_sequence" in item["evidence"]
    ]
    assert len(architecture) == 1
    assert {architecture[0]["book_id"], architecture[0]["compared_book_id"]} == {
        first.book_id,
        second.book_id,
    }
    reloaded = service.current_map(series_id)
    assert reloaded is not None
    persisted = [
        item
        for item in reloaded.findings
        if item["dimension"] == "ARCHITECTURE" and "identical_chapter_sequence" in item["evidence"]
    ]
    assert len(persisted) == 1
    assert persisted[0]["evidence"]["comparison_basis"] == "CURRENT_BOOK_ARCHITECTURE"
    with pytest.raises(SeriesWorkspaceGateError, match="blocking overlap"):
        service.approve_map(series_id, result.map_hash, "Клон нельзя утвердить")


def test_unapproved_book_passport_blocks_writing_gate(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    request = book("Черновой паспорт", 1, "Новая отдельная задача")
    request.lifecycle = "DEFINITION"
    service.add_book(series_id, request)
    result = service.analyze(series_id)
    service.approve_map(series_id, result.map_hash, "Различия проверены")
    with pytest.raises(SeriesWorkspaceGateError, match="Book Passports"):
        service.require_current_map(series_id)


def test_partial_import_fails_closed_and_cannot_be_approved(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    entry = service.add_book(series_id, book("Скан", 1, "Разобрать скан книги"))
    service.import_source(
        series_id,
        entry.book_id,
        SeriesImportRequest(
            filename="scan.pdf",
            content_base64=base64.b64encode(b"not a valid pdf").decode(),
            rights_status="AUTHOR_MANUSCRIPT",
        ),
    )
    result = service.analyze(series_id)
    assert result.status == "BLOCKING"
    assert any(item["dimension"] == "SOURCE_QUALITY" for item in result.findings)
    with pytest.raises(SeriesWorkspaceGateError, match="blocking overlap"):
        service.approve_map(series_id, result.map_hash, "Пусть система проигнорирует ошибку")


def test_services_promotion_preset_is_idempotent_and_keeps_owner_order(tmp_path: Path) -> None:
    service = SeriesWorkspaceService(tmp_path)
    author_id = approved_author(tmp_path)
    first = service.create_services_promotion_preset(
        SeriesPresetRequest(author_profile_id=author_id)
    )
    second = service.create_services_promotion_preset(
        SeriesPresetRequest(author_profile_id=author_id)
    )

    assert first.series_profile_id == second.series_profile_id
    assert first.name == "Секреты продвижения услуг"
    assert [item.title for item in first.books[:4]] == [
        "Как продавать услуги",
        "Секреты продвижения услуг психолога в Яндекс Директ",
        "Как продвигать юридические услуги в Яндекс Директ: Практическое руководство",
        "Как продать онлайн-курсы",
    ]
    assert len(second.books) == 8
    assert second.profile_status == "DRAFT"
    assert [item.origin_kind for item in second.books] == [
        "CURRENT_REWRITTEN",
        "LEGACY_TITLE_ONLY",
        "LEGACY_TITLE_ONLY",
        "LEGACY_TITLE_ONLY",
        "NEW",
        "NEW",
        "NEW",
        "NEW",
    ]
    assert second.books[0].lifecycle == "COMPLETED"
    assert second.books[0].current_corpus_eligible
    assert [item.title for item in second.books[4:]] == [
        "Как продавать услуги компаниям: от первого контакта до договора",
        "Как продвигать местные услуги: клиенты в вашем городе и районе",
        "Как продавать дорогие услуги: доверие, доказательства и выбор исполнителя",
        "Как возвращать клиентов: повторные продажи и рекомендации в услугах",
    ]


def test_legacy_title_only_starts_fresh_and_inherits_only_exact_title(tmp_path: Path) -> None:
    service = SeriesWorkspaceService(tmp_path)
    workspace = service.create_services_promotion_preset(
        SeriesPresetRequest(author_profile_id=approved_author(tmp_path))
    )
    legacy = workspace.books[1]
    old_project_dir = tmp_path / "projects" / legacy.book_id
    legacy_file = old_project_dir / "series-imports" / "old.txt"
    legacy_file.parent.mkdir(parents=True)
    legacy_file.write_text("СТАРЫЙ ТЕКСТ НЕ ДОЛЖЕН ПОПАСТЬ В НОВУЮ КНИГУ", encoding="utf-8")
    engine = create_database(old_project_dir / "project.sqlite")
    try:
        with engine.begin() as connection:
            connection.execute(
                text(
                    "INSERT INTO series_imported_sources(source_id,series_profile_id,book_id,"
                    "filename,format,relative_path,content_hash,rights_status,analysis_status,"
                    "analysis_json,created_at) VALUES (:source,:series,:book,'old.txt','TXT',"
                    "'series-imports/old.txt',:hash,'AUTHOR_MANUSCRIPT','PARSED',:analysis,:created)"
                ),
                {
                    "source": "01JLEGACY0000000000000000",
                    "series": workspace.series_profile_id,
                    "book": legacy.book_id,
                    "hash": "a" * 64,
                    "analysis": json.dumps(
                        {"sample": "СТАРЫЙ ТЕКСТ", "headings": ["Старое оглавление"]}
                    ),
                    "created": "2026-09-13T00:00:00Z",
                },
            )
    finally:
        engine.dispose()

    fresh = service.start_fresh_book(workspace.series_profile_id, legacy.book_id)
    assert fresh.book_id != legacy.book_id
    assert fresh.title == "Секреты продвижения услуг психолога в Яндекс Директ"
    assert fresh.origin_kind == "LEGACY_TITLE_ONLY"
    assert fresh.lifecycle == "DEFINITION"
    assert fresh.imported_sources == []
    reloaded = next(
        item
        for item in service.workspaces()
        if item.series_profile_id == workspace.series_profile_id
    )
    assert reloaded.books[1].book_id == fresh.book_id
    assert reloaded.books[1].lifecycle == "DEFINITION"
    preserved_legacy_file = tmp_path / "library" / legacy.book_id / "series-imports" / "old.txt"
    assert preserved_legacy_file.exists(), "owner file is preserved in the library, not deleted"
    assert not (tmp_path / "projects" / fresh.book_id / "series-imports").exists()

    context = service.generation_context(fresh.book_id)
    assert context is not None
    assert context["allowed_legacy_fields"] == ["title"]
    assert context["legacy_payload_allowed"] is False
    for forbidden in (
        "legacy_manuscript",
        "legacy_outline",
        "legacy_chapters",
        "legacy_examples",
        "legacy_sources",
    ):
        assert context[forbidden] is None
    assert "СТАРЫЙ ТЕКСТ" not in json.dumps(context, ensure_ascii=False)


def test_rewritten_book_is_anti_duplication_corpus_not_a_generation_template(
    tmp_path: Path,
) -> None:
    service = SeriesWorkspaceService(tmp_path)
    workspace = service.create_services_promotion_preset(
        SeriesPresetRequest(author_profile_id=approved_author(tmp_path))
    )
    fresh = service.start_fresh_book(workspace.series_profile_id, workspace.books[1].book_id)
    context = service.generation_context(fresh.book_id)
    assert context is not None
    assert context["current_corpus_usage"] == "ANTI_DUPLICATION_NOT_GENERATION_TEMPLATE"
    assert [item["title"] for item in context["current_corpus"]] == ["Как продавать услуги"]
    assert all("manuscript" not in item for item in context["current_corpus"])


def test_new_planned_book_has_no_legacy_payload_and_preserves_full_title(tmp_path: Path) -> None:
    service = SeriesWorkspaceService(tmp_path)
    workspace = service.create_services_promotion_preset(
        SeriesPresetRequest(author_profile_id=approved_author(tmp_path))
    )
    planned = workspace.books[4]
    assert planned.title == "Как продавать услуги компаниям: от первого контакта до договора"
    assert planned.origin_kind == "NEW"
    assert not planned.legacy_content_allowed
    fresh = service.start_fresh_book(workspace.series_profile_id, planned.book_id)
    assert fresh.title == planned.title
    assert fresh.lifecycle == "DEFINITION"
    assert fresh.imported_sources == []


def test_legacy_title_only_rejects_import_instead_of_guessing_owner_intent(
    tmp_path: Path,
) -> None:
    service = SeriesWorkspaceService(tmp_path)
    workspace = service.create_services_promotion_preset(
        SeriesPresetRequest(author_profile_id=approved_author(tmp_path))
    )
    legacy = workspace.books[2]
    with pytest.raises(SeriesWorkspaceGateError, match="title only"):
        service.import_source(
            workspace.series_profile_id,
            legacy.book_id,
            SeriesImportRequest(
                filename="old.txt",
                content_base64=base64.b64encode(b"old manuscript").decode(),
                rights_status="AUTHOR_MANUSCRIPT",
            ),
        )


def test_import_can_be_explicitly_deleted_and_book_archived(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    entry = service.add_book(series_id, book("Удаляемый источник", 1, "Проверить удаление"))
    imported = service.import_source(
        series_id,
        entry.book_id,
        SeriesImportRequest(
            filename="source.txt",
            content_base64=base64.b64encode("Текст автора".encode()).decode(),
            rights_status="AUTHOR_MANUSCRIPT",
        ),
    )
    source_path = next((tmp_path / "projects" / entry.book_id / "series-imports").iterdir())
    deleted = service.delete_import(series_id, entry.book_id, imported.source_id)
    assert deleted.status == "DELETED"
    assert not source_path.exists()
    archived = service.archive_book(series_id, entry.book_id)
    assert archived.status == "ARCHIVED"
    assert service.books(series_id)[0].status == "ARCHIVED"


def test_selected_series_export_does_not_start_planned_books(tmp_path: Path) -> None:
    service, series_id = new_series(tmp_path)
    service.add_book(series_id, book("Готовая", 1, "Самостоятельная готовая книга"))
    planned = service.add_book(series_id, book("Будущая", 2, "Отдельная будущая книга"))
    result = service.export_series(
        series_id,
        SeriesExportSelection(
            descriptions=True,
            series_and_book_passports=True,
            difference_map=True,
            sources_and_freshness=True,
            next_books_plan=True,
        ),
    )
    assert any(path.endswith("manifest.json") for path in result.files)
    assert any(path.endswith("План-следующих-книг.json") for path in result.files)
    assert service.books(series_id)[1].book_id == planned.book_id
    assert service.books(series_id)[1].status == "IDEA"
